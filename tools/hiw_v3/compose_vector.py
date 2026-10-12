#!/usr/bin/env python3
"""Title + logo on the vector How It Works thumbnails, every option, plus an options deck.

    python3 tools/hiw_v3/vector.py          # draw the art (docs/vector-thumbnails/)
    python3 tools/hiw_v3/compose_vector.py  # thumbnails, review sheet, PowerPoint

Writes exports/stephanie/how-it-works/vector-options/<slug>-<opt>.jpg (1920x1080,
under 400 KB), layers/ for the deck, docs/how-it-works-vector-options.png and
exports/stephanie/how-it-works/how-it-works-thumbnails-vector-options.pptx.
The live thumbnails/ folder and how-it-works-thumbnails.pptx are not touched.
"""
import os, sys

from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import compose as C   # shared title metrics, sprig-free; JPEG and font helpers

ROOT = C.ROOT
ART = os.path.join(ROOT, "docs", "vector-thumbnails")
OUT = os.path.join(C.HIW, "vector-options")
W, H, MX, MY = C.W, C.H, C.MX, C.MY
WHITE = (255, 255, 255)
LOGO_H = 170

TITLES = {
    "what-is-the-soil-food-web": ("What is the Soil Food Web?", ["What is the", "Soil Food Web?"]),
    "nutrient-cycling": ("Nutrient Cycling", ["Nutrient", "Cycling"]),
    "formation-of-structure": ("Formation of Structure", ["Formation of", "Structure"]),
    "weed-suppression": ("Weed Suppression", ["Weed", "Suppression"]),
    "inhibiting-pests-and-diseases": ("Inhibiting Pests & Diseases", ["Inhibiting Pests", "& Diseases"]),
    "sequestering-carbon": ("Sequestering Carbon", ["Sequestering", "Carbon"]),
}
TL, TR, BL, BR = "tl", "tr", "bl", "br"
# option: (title corner or "lm" left-middle, title colour, logo corner, logo colour, short description)
LAYOUT = {
    ("what-is-the-soil-food-web", "A"): (TL, WHITE, BL, WHITE, "Close-up: root tip with bacteria, fungi, a protozoan and a nematode, on Soil Brown"),
    ("what-is-the-soil-food-web", "B"): (TL, WHITE, BR, WHITE, "Bean plant on Food Web Green, magnifier on its root zone"),
    ("nutrient-cycling", "A"): (BL, WHITE, TR, WHITE, "Close-up: protozoan eats bacteria, nutrients flow to the root hair"),
    ("nutrient-cycling", "B"): (TL, C.GREEN, BR, WHITE, "Lettuce on Organic Cream, magnifier on the nutrient loop"),
    ("formation-of-structure", "A"): (TL, WHITE, BL, WHITE, "Soil crumbs glued by bacteria, woven by fungal threads, on Soil Brown"),
    ("formation-of-structure", "B"): (TL, WHITE, BL, WHITE, "The same crumbs as a specimen on Food Web Green"),
    ("weed-suppression", "A"): (TL, C.GREEN, TR, WHITE, "Weeds in bacterial soil vs corn in fungal soil; cream and green sky like the reference"),
    ("weed-suppression", "B"): (TL, WHITE, TR, WHITE, "Same, Food Web Green sky"),
    ("weed-suppression", "C"): (TL, C.GREEN, TR, C.GREEN, "Same, Organic Cream sky"),
    ("inhibiting-pests-and-diseases", "A"): (TL, WHITE, BL, WHITE, "Close-up: root sleeved in beneficials, pathogen kept out, on Soil Brown"),
    ("inhibiting-pests-and-diseases", "B"): (TL, WHITE, BL, WHITE, "The same as a specimen on Food Web Green"),
    ("inhibiting-pests-and-diseases", "C"): (TL, C.GREEN, BR, WHITE, "Tomato on Organic Cream, magnifier on its root zone"),
    ("sequestering-carbon", "A"): (TR, WHITE, BL, WHITE, "Oak on Food Web Green: carbon flows down roots and fungal threads into dark soil clumps"),
    ("sequestering-carbon", "B"): (TR, C.GREEN, BL, WHITE, "Same, Organic Cream sky"),
}


def logo(colour):
    src = Image.open(os.path.join(ROOT, "img", "brand", "sfw-school-logo-white.png")).convert("RGBA")
    im = Image.new("RGBA", src.size, colour + (255,))
    im.putalpha(src.getchannel("A"))
    return im.resize((round(src.width * LOGO_H / src.height), LOGO_H), Image.LANCZOS)


LINES_OVERRIDE = {("what-is-the-soil-food-web", "A"): ["What is the", "Soil Food", "Web?"]}


def title_box(lines, corner):
    f = C.font()
    lh = round(C.TITLE_PX * C.LEADING)
    widths = [f.getlength(t) for t in lines]
    bw, bh = max(widths), lh * (len(lines) - 1) + C.TITLE_PX
    x = W - MX - bw if corner in (TR, BR) else MX
    y = {TL: MY, TR: MY, BL: H - MY - bh, BR: H - MY - bh, "lm": (H - bh) / 2}[corner]
    align = "right" if corner in (TR, BR) else "left"
    return x, y, bw, bh, lh, widths, align


def logo_pos(lg, corner):
    return (W - MX - lg.width if corner in (TR, BR) else MX, MY if corner in (TL, TR) else H - MY - lg.height)


def build(slug, opt):
    tcorner, tcol, lcorner, lcol, _ = LAYOUT[(slug, opt)]
    art = Image.open(os.path.join(ART, "%s-%s.png" % (slug, opt))).convert("RGB")
    lg = logo(lcol)
    im = art.copy()
    d = ImageDraw.Draw(im)
    lines = LINES_OVERRIDE.get((slug, opt), TITLES[slug][1])
    x, y, bw, bh, lh, widths, align = title_box(lines, tcorner)
    for i, (t, w) in enumerate(zip(lines, widths)):
        lx = x + bw - w if align == "right" else x
        d.text((lx, y + i * lh - C.TITLE_PX * 0.21), t, font=C.font(), fill=tcol)
    lp = logo_pos(lg, lcorner)
    im.paste(lg, lp, lg)
    return art, lg, lp, im, (x, y, bw, bh, lh, align)


def main():
    os.makedirs(os.path.join(OUT, "layers"), exist_ok=True)
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    px = 12192000 / W
    done = []
    for slug in TITLES:
        for opt in sorted(o for s, o in LAYOUT if s == slug):
            art, lg, lp, im, (x, y, bw, bh, lh, align) = build(slug, opt)
            name = "%s-%s" % (slug, opt)
            q, n = C.save_jpg(im, os.path.join(OUT, name + ".jpg"))
            ap, lgp = os.path.join(OUT, "layers", name + "-art.png"), os.path.join(OUT, "layers", name + "-logo.png")
            art.save(ap, optimize=True)
            lg.save(lgp)
            print("%-36s q%d %d KB" % (name, q, n // 1024))
            done.append((slug, opt, im))

            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(ap, 0, 0, prs.slide_width, prs.slide_height).name = "Artwork"
            s.shapes.add_picture(lgp, Emu(round(lp[0] * px)), Emu(round(lp[1] * px)),
                                 Emu(round(lg.width * px)), Emu(round(lg.height * px))).name = "Logo"
            pad = 24
            tb = s.shapes.add_textbox(Emu(round((x - pad) * px)), Emu(round((y - C.TITLE_PX * 0.3) * px)),
                                      Emu(round((bw + 2 * pad) * px)), Emu(round((bh + C.TITLE_PX * 0.5) * px)))
            tb.name = "Title"
            tf = tb.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(tf, m, 0)
            tcol = LAYOUT[(slug, opt)][1]
            for i, line in enumerate(LINES_OVERRIDE.get((slug, opt), TITLES[slug][1])):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.RIGHT if align == "right" else PP_ALIGN.LEFT
                p.line_spacing = Pt(lh * 0.5)
                r = p.add_run()
                r.text = line
                r.font.name = "Montserrat"
                r.font.bold = True
                r.font.size = Pt(C.TITLE_PX * 0.5)
                r.font.color.rgb = RGBColor(*tcol)
            s.notes_slide.notes_text_frame.text = "%s, option %s: %s. Vector illustration (no AI image generation)." % (
                TITLES[slug][0], opt, LAYOUT[(slug, opt)][4])
    prs.save(os.path.join(C.HIW, "how-it-works-thumbnails-vector-options.pptx"))

    # review sheet: one row per video, options side by side, plus 168 px copies
    tw, th, g = 620, 349, 24
    f = ImageFont.truetype(C.FONT, 24)
    rows = list(TITLES)
    sheet = Image.new("RGB", (g + 3 * (tw + g) + 3 * (168 + 12) + g, len(rows) * (th + 54) + g), (244, 241, 234))
    d = ImageDraw.Draw(sheet)
    for r, slug in enumerate(rows):
        y = g + r * (th + 54)
        d.text((g, y), "%d  %s" % (r + 1, TITLES[slug][0]), font=f, fill=(35, 31, 29))
        opts = [(o, im) for s_, o, im in done if s_ == slug]
        for c, (o, im) in enumerate(opts):
            x = g + c * (tw + g)
            sheet.paste(im.resize((tw, th), Image.LANCZOS), (x, y + 32))
            d.text((x + 8, y + 36), o, font=f, fill=(255, 255, 255), stroke_width=3, stroke_fill=(35, 31, 29))
        for c, (o, im) in enumerate(opts):
            x = g + 3 * (tw + g) + c * (168 + 12)
            sheet.paste(im.resize((168, 95), Image.LANCZOS), (x, y + 32))
    sheet.save(os.path.join(ROOT, "docs", "how-it-works-vector-options.png"), optimize=True)


if __name__ == "__main__":
    main()
