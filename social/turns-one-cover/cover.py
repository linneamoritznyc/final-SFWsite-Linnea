"""Cover slide for the one-year anniversary carousel: a soft halo at the centre, community photos and
flower cutouts orbiting it, EB Garamond type. Three compositions.

  python3 social/turns-one-cover/cover.py            -> cover-1/ .. cover-3/ (post.png, layers/, pptx) + covers-compare.png
  python3 social/turns-one-cover/cover.py 2          -> only composition 2
Light comes from the top left: every shadow falls down and to the right.
"""
import os, sys, json, math, random, subprocess
import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
NEW = os.path.join(ROOT, "img", "new-2026-10")
FLOW = os.path.join(ROOT, "social", "cutouts-flowers")
ORCHID = os.path.join(ROOT, "social", "open-house-2026-10-16", "cutouts", "orchid-costa-rica.webp")
W, H = 1080, 1350
DEEP, GREEN, GOLD, PAPER = "#22371F", "#156826", "#C9A227", "#FBF8F2"
SHADOW = (12, 18)  # light from the top left


def rgb(h): return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))


PHOTOS = {  # name: (file, focus x, focus y)
    "high-five-compost": (os.path.join(NEW, "students-high-five-at-compost.jpg"), .45, .2),
    "microscope-santa-fe": (os.path.join(NEW, "santa-fe-2026-21-at-the-microscope.jpg"), .45, .3),
    "laughing-at-microscope": (os.path.join(NEW, "community-laughing-at-microscope.jpg"), .45, .3),
    "hands-in-compost": (os.path.join(NEW, "costa-rica-2024-07-11-hands-in-the-compost.jpg"), .5, .55),
    "poppy-field-group": (os.path.join(NEW, "building-community-poppy-field-group.jpg"), .5, .55),
    "soil-in-hand": (os.path.join(NEW, "costa-rica-2024-07-10-moisture-squeeze-test.jpg"), .6, .3),
    "mentor-microscopes": (os.path.join(NEW, "santa-fe-2026-23-mentor-with-microscopes.jpg"), .4, .35),
}
FLOWERS = {
    "hibiscus": os.path.join(FLOW, "costa-rica-hibiscus.png"), "marigold": os.path.join(FLOW, "tamil-nadu-marigold.png"),
    "orchid": ORCHID, "foxglove": os.path.join(FLOW, "norfolk-foxglove.png"),
    "jasmine": os.path.join(FLOW, "tamil-nadu-jasmine.png"), "poppy": os.path.join(FLOW, "norfolk-common-poppy.png"),
}


def cover(path, w, h, fx, fy):
    im = Image.open(path).convert("RGB"); s = max(w / im.width, h / im.height)
    im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    x = min(max(0, round(im.width * fx - w / 2)), im.width - w); y = min(max(0, round(im.height * fy - h / 2)), im.height - h)
    return im.crop((x, y, x + w, y + h))


def mask(shape, w, h):
    m = Image.new("L", (w * 2, h * 2), 0); d = ImageDraw.Draw(m)
    if shape == "circle": d.ellipse((0, 0, w * 2 - 1, h * 2 - 1), fill=255)
    elif shape == "arch": d.rectangle((0, w, w * 2, h * 2), fill=255); d.ellipse((0, 0, w * 2 - 1, w * 2 - 1), fill=255)
    else: d.rounded_rectangle((0, 0, w * 2 - 1, h * 2 - 1), radius=28, fill=255)
    return m.resize((w, h), Image.LANCZOS)


def photo(name, shape, w, h):
    f, fx, fy = PHOTOS[name]
    im = cover(f, w, h, fx, fy).convert("RGBA"); im.putalpha(mask(shape, w, h)); return im


def flower(name, h, rot=0):
    im = Image.open(FLOWERS[name]).convert("RGBA")
    im = im.crop(im.getchannel("A").point(lambda a: 255 if a > 10 else 0).getbbox())
    im = im.resize((round(im.width * h / im.height), h), Image.LANCZOS)
    return im.rotate(rot, expand=True, resample=Image.BICUBIC) if rot else im


def with_shadow(im, blur=16, alpha=.30, off=SHADOW):
    pad = blur * 3 + max(off)
    out = Image.new("RGBA", (im.width + 2 * pad, im.height + 2 * pad), (0, 0, 0, 0))
    sh = Image.new("RGBA", im.size, (60, 42, 30, 255)); sh.putalpha(im.getchannel("A").point(lambda v: int(v * alpha)))
    layer = Image.new("RGBA", out.size, (0, 0, 0, 0)); layer.alpha_composite(sh, (pad + off[0], pad + off[1]))
    out.alpha_composite(layer.filter(ImageFilter.GaussianBlur(blur))); out.alpha_composite(im, (pad, pad))
    return out, pad


def paper_bg(seed):
    rng = np.random.default_rng(seed)
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    base = np.array(rgb(PAPER), np.float32)
    vign = (np.hypot((xx - W * .42) / W, (yy - H * .4) / H) * 14)[..., None]       # slightly warmer away from the light
    a = base - vign * np.array([0.6, 0.9, 1.6]) + rng.normal(0, 2.2, (H, W, 1))
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))


def halo(w, h, shape="arch"):
    """Soft glowing halo: a warm light with a feathered edge, like the arch in the painting."""
    big = Image.new("RGBA", (w + 240, h + 240), (0, 0, 0, 0))
    m = mask(shape, w, h)
    glow = Image.new("L", big.size, 0); glow.paste(m, (120, 120)); glow = glow.filter(ImageFilter.GaussianBlur(60))
    core = Image.new("L", big.size, 0); core.paste(m, (120, 120)); core = core.filter(ImageFilter.GaussianBlur(14))
    yy, xx = np.mgrid[0:big.height, 0:big.width].astype(np.float32)
    rad = np.clip(1 - np.hypot((xx - big.width * .45) / big.width, (yy - big.height * .35) / big.height) * 1.3, 0, 1)
    col = np.zeros((big.height, big.width, 4), np.float32)
    col[..., :3] = np.array([252, 226, 178], np.float32) * (1 - rad[..., None]) + np.array([255, 246, 226], np.float32) * rad[..., None]
    col[..., 3] = np.maximum(np.asarray(glow, np.float32) * .85, np.asarray(core, np.float32) * 1.0)
    return Image.fromarray(col.astype(np.uint8))


def candle(h=230):
    w = 36; im = Image.new("RGBA", (w + 60, h + 90), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    x0 = 30; d.rounded_rectangle((x0, 90, x0 + w, 90 + h), radius=6, fill=rgb("#F3E7C9") + (255,))
    for y in range(110, 90 + h, 30): d.polygon([(x0, y), (x0 + w, y - 14), (x0 + w, y - 4), (x0, y + 10)], fill=rgb(GOLD) + (255,))
    d.line((x0 + w // 2, 76, x0 + w // 2, 92), fill=(70, 50, 35, 255), width=3)
    f = Image.new("RGBA", im.size, (0, 0, 0, 0)); fd = ImageDraw.Draw(f)
    fd.ellipse((x0 + w // 2 - 13, 20, x0 + w // 2 + 13, 80), fill=(255, 176, 60, 255)); fd.ellipse((x0 + w // 2 - 6, 42, x0 + w // 2 + 6, 78), fill=(255, 238, 190, 255))
    im.alpha_composite(f); return im


def candle_glow():
    g = Image.new("RGBA", (260, 260), (0, 0, 0, 0)); ImageDraw.Draw(g).ellipse((70, 60, 190, 200), fill=(255, 200, 100, 150))
    return g.filter(ImageFilter.GaussianBlur(30))


def confetti(seed, n, box):
    rnd = random.Random(seed); pieces = []
    cols = [GOLD, GREEN, "#8E6BB0", "#E07A5F", "#F2C14E"]
    for i in range(n):
        w, h = rnd.choice([(26, 12), (18, 18), (30, 10), (14, 14)])
        im = Image.new("RGBA", (w, h), rgb(rnd.choice(cols)) + (255,))
        if w == h and rnd.random() < .5: im.putalpha(mask("circle", w, h))
        im = im.rotate(rnd.uniform(0, 180), expand=True, resample=Image.BICUBIC)
        pieces.append((im, rnd.randint(box[0], box[2]), rnd.randint(box[1], box[3])))
    return pieces


class Cover:
    def __init__(self, key):
        self.key, self.dir = key, os.path.join(HERE, f"cover-{key}")
        self.lay = os.path.join(self.dir, "layers"); os.makedirs(self.lay, exist_ok=True)
        for f in os.listdir(self.lay): os.remove(os.path.join(self.lay, f))
        self.items = []
        self.add("background", paper_bg(int(key) * 7), 0, 0, shadow=False, jpg=True)

    def add(self, name, im, cx, cy, shadow=True, jpg=False, tl=True, blur=16, alpha=.30):
        """Place im with its top-left (or centre if tl=False) at cx, cy; shadow is its own layer under it."""
        x, y = (cx, cy) if tl else (round(cx - im.width / 2), round(cy - im.height / 2))
        if shadow:
            sh, pad = with_shadow(im, blur, alpha)
            only = sh.copy(); hole = Image.new("RGBA", sh.size, (0, 0, 0, 0)); hole.alpha_composite(im, (pad, pad))
            a = np.asarray(only, np.float32); b = np.asarray(hole, np.float32)[..., 3:4] / 255
            a[..., 3:4] *= (1 - b); only = Image.fromarray(a.astype(np.uint8))
            self._save(name + "-shadow", only, x - pad, y - pad)
        self._save(name, im, x, y, jpg)

    def _save(self, name, im, x, y, jpg=False):
        box = (max(0, -x), max(0, -y), min(im.width, W - x), min(im.height, H - y)); im = im.crop(box); x += box[0]; y += box[1]
        fn = f"{len(self.items) + 1:02d}-{name}." + ("jpg" if jpg else "png")
        (im.convert("RGB").save(os.path.join(self.lay, fn), quality=90) if jpg else im.save(os.path.join(self.lay, fn), optimize=True))
        self.items.append(dict(type="image", name=name, file="layers/" + fn, x=x, y=y, w=im.width, h=im.height))

    def text(self, name, t, x, y, w, h, size, italic, color, align="left", weight=500):
        self.items.append(dict(type="text", name=name, text=t, x=x, y=y, w=w, h=h, size=size, italic=italic, color=color, align=align, weight=weight))

    def logo(self, x, y, w=120):
        l = Image.open(os.path.join(ROOT, "img", "sfwlogo-240.png")).convert("RGBA")
        self._save("logo", l.resize((w, round(w * l.height / l.width)), Image.LANCZOS), x, y)

    def write(self):
        fonts = "file://" + os.path.join(ROOT, "fonts")
        css = f"""@font-face{{font-family:"EB Garamond";src:url("{fonts}/eb-garamond-latin-500-italic.woff2");font-style:italic;font-weight:500}}
@font-face{{font-family:"EB Garamond";src:url("{fonts}/eb-garamond-latin-400-normal.woff2");font-style:normal;font-weight:400}}
*{{margin:0}}body{{background:#555;padding:40px}}.slide{{position:relative;width:1080px;height:1350px;overflow:hidden}}
.slide img{{position:absolute;display:block}}.slide p{{position:absolute;font-family:"EB Garamond";line-height:1.02}}"""
        out = [f"<!doctype html><html><head><meta charset='utf-8'><title>Cover {self.key}</title><style>{css}</style></head><body><section class='slide'>"]
        for it in self.items:
            if it["type"] == "image":
                out.append(f"<img src='{it['file']}' style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;height:{it['h']}px' alt=''>")
            else:
                out.append(f"<p style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;font-size:{it['size']}px;font-weight:{it['weight']};"
                           f"font-style:{'italic' if it['italic'] else 'normal'};color:{it['color']};text-align:{it['align']}'>{it['text']}</p>")
        out.append("</section></body></html>")
        open(os.path.join(self.dir, "slides.html"), "w").write("\n".join(out))
        json.dump(self.items, open(os.path.join(self.dir, "layout.json"), "w"), indent=1)
        subprocess.run(["node", os.path.join(ROOT, "social", "turns-one-2026-10-16", "render.cjs"), self.dir], check=True)
        os.replace(os.path.join(self.dir, "slide-1.png"), os.path.join(self.dir, "post.png"))
        if os.path.exists(os.path.join(self.dir, "contact-sheet.png")): os.remove(os.path.join(self.dir, "contact-sheet.png"))
        from pptx import Presentation
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        E = 914400 / 108
        prs = Presentation(); prs.slide_width, prs.slide_height = Emu(round(W * E)), Emu(round(H * E))
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for it in self.items:
            if it["type"] == "image":
                sl.shapes.add_picture(os.path.join(self.dir, it["file"]), *[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]).name = it["name"]
            else:
                tb = sl.shapes.add_textbox(*[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]); tb.name = it["name"]
                tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if it["align"] == "center" else PP_ALIGN.LEFT
                r = p.add_run(); r.text = it["text"]; f = r.font
                f.name = "EB Garamond"; f.size = Pt(round(it["size"] * 72 / 108, 1)); f.italic = it["italic"]
                f.color.rgb = RGBColor(*rgb(it["color"]))
        prs.save(os.path.join(self.dir, f"turns-one-{os.path.basename(self.dir)}-canva.pptx"))


HEAD, SUB = "Today we turn 1", "One year as a 501(c)(3) nonprofit."


def build1():
    """1. The arch: a tall glowing arch in the centre holds the words; photos orbit it clockwise."""
    c = Cover("1")
    c.add("halo", halo(560, 760, "arch"), 540, 640, shadow=False, tl=False)
    c.add("photo-poppy-field-group", photo("poppy-field-group", "arch", 300, 380), 70, 150)
    c.add("flower-foxglove", flower("foxglove", 330, 8), 300, 120)
    c.add("photo-high-five-compost", photo("high-five-compost", "rect", 360, 250), 660, 110)
    c.add("flower-orchid", flower("orchid", 190, -10), 640, 380, tl=False)
    c.add("photo-microscope-santa-fe", photo("microscope-santa-fe", "circle", 270, 270), 770, 560)
    c.add("photo-hands-in-compost", photo("hands-in-compost", "arch", 250, 330), 780, 900)
    c.add("flower-hibiscus", flower("hibiscus", 210, 10), 760, 1220, tl=False)
    c.add("photo-laughing-at-microscope", photo("laughing-at-microscope", "rect", 380, 260), 330, 1010)
    c.add("flower-marigold", flower("marigold", 230, -6), 280, 1000, tl=False)
    c.add("photo-soil-in-hand", photo("soil-in-hand", "circle", 250, 250), 60, 640)
    c.add("flower-jasmine", flower("jasmine", 170, 0), 130, 600, tl=False)
    c.add("flower-poppy", flower("poppy", 260, -8), 160, 1180, tl=False)
    c.add("candle-glow", candle_glow(), 540, 400 - 90, shadow=False, tl=False)
    c.add("candle", candle(200), 540, 400, shadow=True, tl=False, blur=10, alpha=.22)
    for i, (im, x, y) in enumerate(confetti(11, 9, (360, 300, 720, 900))):
        if 300 < x < 780 and 520 < y < 900: continue
        c.add(f"confetti-{i + 1}", im, x, y, blur=4, alpha=.25)
    c.text("headline", HEAD, 340, 600, 400, 300, 104, True, DEEP, "center")
    c.text("small-line", SUB, 330, 830, 420, 80, 32, False, GREEN, "center", 400)
    c.logo(W - 60 - 110, H - 60 - 95, 110)
    c.write()


def build2():
    """2. The round halo: one big circular photo glowing at the centre, smaller photos tucked around it, words below."""
    c = Cover("2")
    c.add("halo", halo(640, 640, "circle"), 560, 520, shadow=False, tl=False)
    c.add("photo-microscope-santa-fe", photo("microscope-santa-fe", "rect", 300, 220), 60, 170)
    c.add("photo-poppy-field-group", photo("poppy-field-group", "rect", 330, 230), 700, 120)
    c.add("flower-foxglove", flower("foxglove", 360, -8), 120, 470, tl=False)
    c.add("photo-high-five-compost", photo("high-five-compost", "circle", 470, 470), 325, 285)
    c.add("flower-orchid", flower("orchid", 170, 12), 860, 760, tl=False)
    c.add("photo-hands-in-compost", photo("hands-in-compost", "arch", 220, 290), 790, 470)
    c.add("photo-laughing-at-microscope", photo("laughing-at-microscope", "circle", 200, 200), 140, 560)
    c.add("flower-hibiscus", flower("hibiscus", 180, -12), 330, 735, tl=False)
    c.add("flower-marigold", flower("marigold", 170, 8), 990, 380, tl=False)
    c.add("flower-jasmine", flower("jasmine", 150), 640, 250, tl=False)
    c.add("candle-glow", candle_glow(), 560, 220 - 90, shadow=False, tl=False)
    c.add("candle", candle(170), 560, 220, tl=False, blur=10, alpha=.22)
    for i, (im, x, y) in enumerate(confetti(23, 10, (80, 60, 1000, 900))):
        if 260 < x < 860 and 240 < y < 820: continue
        c.add(f"confetti-{i + 1}", im, x, y, blur=4, alpha=.25)
    c.text("headline", HEAD, 90, 900, 900, 150, 128, True, DEEP, "center")
    c.text("small-line", SUB, 90, 1060, 900, 60, 38, False, GREEN, "center", 400)
    c.add("flower-poppy", flower("poppy", 230, -6), 140, 1220, tl=False)
    c.logo(W // 2 - 55, H - 50 - 95, 110)
    c.write()


def build3():
    """3. Off-centre: a halo arch on the right holds a tall photo; words top left; smaller photos loop down the left."""
    c = Cover("3")
    c.add("halo", halo(520, 820, "arch"), 700, 760, shadow=False, tl=False)
    c.add("photo-mentor-microscopes", photo("mentor-microscopes", "arch", 400, 620), 500, 430)
    c.add("flower-orchid", flower("orchid", 200, -14), 520, 470, tl=False)
    c.add("flower-foxglove", flower("foxglove", 380, 6), 960, 560, tl=False)
    c.add("photo-high-five-compost", photo("high-five-compost", "rect", 400, 270), 70, 520)
    c.add("photo-hands-in-compost", photo("hands-in-compost", "circle", 230, 230), 290, 830)
    c.add("photo-poppy-field-group", photo("poppy-field-group", "rect", 330, 230), 60, 860)
    c.add("photo-laughing-at-microscope", photo("laughing-at-microscope", "arch", 260, 300), 640, 1000)
    c.add("flower-hibiscus", flower("hibiscus", 190, 8), 980, 1200, tl=False)
    c.add("flower-marigold", flower("marigold", 190, -6), 290, 1190, tl=False)
    c.add("flower-jasmine", flower("jasmine", 150), 520, 870, tl=False)
    c.add("flower-poppy", flower("poppy", 240, 6), 545, 1230, tl=False)
    c.add("candle-glow", candle_glow(), 900, 330 - 90, shadow=False, tl=False)
    c.add("candle", candle(190), 900, 330, tl=False, blur=10, alpha=.22)
    for i, (im, x, y) in enumerate(confetti(37, 8, (520, 60, 1020, 380))):
        c.add(f"confetti-{i + 1}", im, x, y, blur=4, alpha=.25)
    c.text("headline", HEAD, 70, 120, 520, 300, 120, True, DEEP, "left")
    c.text("small-line", SUB, 74, 400, 520, 60, 34, False, GREEN, "left", 400)
    c.logo(70, H - 60 - 95, 110)
    c.write()


def compare(keys=("1", "2", "3")):
    tw, th, g = 480, 600, 30
    cmp_ = Image.new("RGB", (g + len(keys) * (tw + g), th + 90), (31, 31, 31)); dr = ImageDraw.Draw(cmp_)
    f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 26)
    names = {"1": "1  The arch", "2": "2  The round halo", "3": "3  Off-centre"}
    for i, k in enumerate(keys):
        x = g + i * (tw + g); dr.text((x, 18), names[k], font=f, fill=(240, 240, 240))
        cmp_.paste(Image.open(os.path.join(HERE, f"cover-{k}", "post.png")).convert("RGB").resize((tw, th), Image.LANCZOS), (x, 60))
    cmp_.save(os.path.join(HERE, "covers-compare.png"))


if __name__ == "__main__":
    which = sys.argv[1:] or ["1", "2", "3"]
    for k in which: {"1": build1, "2": build2, "3": build3}[k]()
    compare(); print("done")
