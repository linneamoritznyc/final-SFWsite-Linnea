import json, re, os, subprocess
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
HERE = os.path.dirname(os.path.abspath(__file__)); EW = os.path.dirname(HERE); os.chdir(HERE)
OUT = '/home/user/final-sfwsite-linnea/social/earthworm-day-2026-10-21/21-10-2026-wed-ig-story-earthworm.pptx'
E = lambda px: Emu(int(round(px * 9525)))
INK = RGBColor(0x4F, 0x34, 0x33)
NOTES = {'frame1': 'Source: https://www.earthwormsoc.org.uk/worldwormday  (Fonts: Montserrat, free on Google Fonts and in Canva.)', 'frame2': 'Sources: https://nautil.us/why-darwin-admired-the-humble-earthworm-361515 ; http://darwin-online.org.uk/', 'frame3': 'Sources: https://www.echocommunity.org/en/resources/58d05ca9-030c-42fa-98fa-746d4dc9bc2f ; https://crsa.cnr.berkeley.edu/wp-content/uploads/2021/08/Soil-Biology-Primer.pdf', 'frame4': 'Source: https://extension.colostate.edu/resource/earthworms  (Leave the empty space under the jars for the poll sticker.)'}
specs = json.load(open('layers/specs.json'))
LINES = json.loads(subprocess.run(['node', f'{EW}/lines.js'] + [s['html'] for s in specs], capture_output=True, text=True, check=True).stdout)
prs = Presentation(); prs.slide_width, prs.slide_height = E(1080), E(1920)
for spec in specs:
    rendered = LINES[spec['html']]
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(255, 255, 255)
    for name, path, x, y, w, h in spec['pics']:
        im = Image.open(path); bb = im.getbbox(); s = w / im.width
        c = im.crop(bb); p = path.replace('.png', '-crop.png'); c.save(p)
        sl.shapes.add_picture(p, E(x + bb[0] * s), E(y + bb[1] * s), E(c.width * s), E(c.height * s)).name = name
    for n, (p0, p1, bend) in enumerate(spec['arrows']):
        mx, my = (p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2
        dx, dy = p1[0] - p0[0], p1[1] - p0[1]
        cx, cy = mx - dy * bend, my + dx * bend
        pts = [((1-t)**2*p0[0] + 2*(1-t)*t*cx + t*t*p1[0], (1-t)**2*p0[1] + 2*(1-t)*t*cy + t*t*p1[1]) for t in [k/24 for k in range(25)]]
        fb = sl.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
        sh = fb.convert_to_shape(); sh.name = f'Arrow {n+1}'
        sh.fill.background(); sh.line.color.rgb = INK; sh.line.width = Pt(2)
        etree.SubElement(sh.line._get_or_add_ln(), qn('a:tailEnd'), type='triangle', w='med', len='med')
    for ti, t in enumerate(spec['texts']):
        t.setdefault('lh', 1.25); t.setdefault('weight', 500); t.setdefault('family', 'body'); t.setdefault('color', '#5b4c46'); t.setdefault('align', 'left')
        bolds = re.findall(r'<b>(.*?)</b>', t['text']); itals = re.findall(r'<i>(.*?)</i>', t['text'])
        lines = rendered[ti]; pad = 80
        bx = {'left': t['x'], 'center': t['x'] - pad / 2, 'right': t['x'] - pad}[t['align']]
        tb = sl.shapes.add_textbox(E(bx), E(t['y'] - 4), E(t['w'] + pad), E(t['size'] * t['lh'] * len(lines) + 10))
        tb.name = re.sub('<[^>]+>', '', t['text'])[:40]
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fam = 'Fredoka' if t['family'] == 'title' else 'Montserrat'
        al = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[t['align']]
        marks = [(b, 'b') for b in bolds] + [(i, 'i') for i in itals]
        for li, l in enumerate(lines):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            p.line_spacing = Pt(t['size'] * t['lh'] * 0.75); p.alignment = al  # exact, matches the PNG
            segs = [(l, '')]
            for phrase, kind in marks:
                for word in phrase.split(' '):  # phrases can wrap across lines: mark word by word
                    new = []
                    for txt, k in segs:
                        if k or word not in txt: new.append((txt, k)); continue
                        a, _, b = txt.partition(word); new += [(a, ''), (word, kind), (b, '')]
                    segs = new
            for txt, k in segs:
                if not txt: continue
                r = p.add_run(); r.text = txt
                r.font.size = Pt(t['size'] * 0.75); r.font.name = fam
                r.font.bold = (t['weight'] >= 600) or k == 'b'; r.font.italic = k == 'i'
                r.font.color.rgb = RGBColor.from_string(t['color'].lstrip('#'))
    lg = Image.open(f'{EW}/logo.png')
    sl.shapes.add_picture(f'{EW}/logo.png', E(890), E(1560), E(150), E(150 * lg.height / lg.width)).name = 'SFW logo'
    sl.notes_slide.notes_text_frame.text = NOTES[spec['name']]
prs.save(OUT); print(OUT)
