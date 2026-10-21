"""World Earthworm Day 2026-10-21: carousel (6 x 1080x1350) and story (3 x 1080x1920).
Each cutout is baked with its own soft drop shadow into a layer PNG; every layer and text box is a separate
object in the .pptx. Run: python3 build_worm_holes.py"""
import os, math, random
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
ONE = os.path.join(HERE, '..', 'turns-one-2026-10-16', 'layers')
CUT = os.path.join(HERE, 'cutouts')
FLOW = os.path.join(HERE, '..', 'cutouts-flowers')
FONTS = '/root/.fonts'
HEAD = os.path.join(FONTS, 'EBGaramond-MediumItalic.ttf')
NUM = os.path.join(FONTS, 'EBGaramond-SemiBold.ttf')
SMALL = os.path.join(FONTS, 'Montserrat-SemiBold.ttf')
CREAM = (244, 241, 234); DEEP = (34, 55, 31); FWG = (21, 104, 38); SOIL = (79, 52, 51)
SAGE = (205, 224, 196)
EMU = 9525  # 96 dpi

def src(name):
    for d in (CUT, ONE, FLOW):
        p = os.path.join(d, name)
        if os.path.exists(p): return p
    raise FileNotFoundError(name)

class Slide:
    def __init__(self, w, h, notes=''):
        self.w, self.h, self.items, self.notes = w, h, [], notes
        self.light = (10, 12)
    def bg(self, name):
        self.items.append(('bg', name))
    def cut(self, name, x, y, w=None, h=None, shadow=True, alpha=1.0, label=None):
        im = name if isinstance(name, Image.Image) else Image.open(src(name)).convert('RGBA')
        if w and not h: h = round(im.height * w / im.width)
        if h and not w: w = round(im.width * h / im.height)
        if not w: w, h = im.size
        im = im.resize((w, h), Image.LANCZOS)
        if alpha < 1: im.putalpha(im.split()[3].point(lambda v: int(v * alpha)))
        self.items.append(('img', im, x, y, shadow, label or (name if isinstance(name, str) else 'shape')))
    def text(self, txt, x, y, w, size, font='head', color=CREAM, lh=1.06, align='left', label='text'):
        self.items.append(('text', txt, x, y, w, size, font, color, lh, align, label))

def with_shadow(im, light, blur=11, opacity=0.4):
    dx, dy = light; pad = blur * 3 + max(abs(dx), abs(dy))
    W, H = im.width + 2 * pad, im.height + 2 * pad
    a = im.split()[3]
    sh = Image.new('L', (W, H), 0); sh.paste(a, (pad + dx, pad + dy))
    sh = sh.filter(ImageFilter.GaussianBlur(blur)).point(lambda v: int(v * opacity))
    out = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    out.putalpha(sh)
    shadow_rgb = Image.new('RGBA', (W, H), (20, 14, 10, 0)); shadow_rgb.putalpha(sh)
    out = Image.alpha_composite(Image.new('RGBA', (W, H), (0, 0, 0, 0)), shadow_rgb)
    out.alpha_composite(im, (pad, pad))
    return out, pad

FONTMAP = {'head': HEAD, 'num': NUM, 'small': SMALL}
PPTFONT = {'head': ('EB Garamond Medium', True, False), 'num': ('EB Garamond SemiBold', False, False), 'small': ('Montserrat SemiBold', False, False)}

def wrap(txt, font, w):
    lines = []
    for para in txt.split('\n'):
        cur = ''
        for word in para.split(' '):
            t = (cur + ' ' + word).strip()
            if font.getlength(t) <= w: cur = t
            else: lines.append(cur); cur = word
        lines.append(cur)
    return lines

def ribbon_layer(points, w, h, width=46):
    paths = points if isinstance(points[0][0], (list, tuple)) else [points]
    """Earthworm burrow drawn as a cut-paper tunnel: dark lining, soil-brown core, fine paper grain."""
    S = 2
    big = Image.new('RGBA', (w * S, h * S), (0, 0, 0, 0)); d = ImageDraw.Draw(big)
    for wd, col in ((width + 14, (52, 33, 30, 255)), (width, (96, 66, 58, 255)), (width - 22, (79, 52, 51, 255))):
        for pts in paths:
            P = [(x * S, y * S) for x, y in pts]
            d.line(P, fill=col, width=wd * S, joint='curve')
            r = wd * S // 2
            for x, y in P: d.ellipse([x - r, y - r, x + r, y + r], fill=col)
    im = big.resize((w, h), Image.LANCZOS)
    rng = np.random.default_rng(3); a = np.asarray(im).astype(float)
    n = rng.normal(0, 6, a.shape[:2])[..., None]; a[..., :3] = np.clip(a[..., :3] + n, 0, 255)
    im = Image.fromarray(a.astype(np.uint8), 'RGBA'); bb = im.getbbox()
    return im.crop(bb), bb[0], bb[1]

def render(slide, out_png, pptx_slide, prs):
    canvas = Image.new('RGBA', (slide.w, slide.h), CREAM + (255,))
    shapes = pptx_slide.shapes
    tmp = os.path.join(HERE, 'build-tmp'); os.makedirs(tmp, exist_ok=True)
    n = 0
    for it in slide.items:
        n += 1
        if it[0] == 'bg':
            im = Image.open(src(it[1])).convert('RGB').resize((slide.w, slide.h), Image.LANCZOS)
            canvas.paste(im, (0, 0)); p = os.path.join(tmp, f'{os.path.basename(out_png)[:-4]}-{n:02d}-background.jpg'); im.save(p, quality=92)
            shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height).name = 'background'
        elif it[0] == 'img':
            _, im, x, y, shadow, label = it
            lay, pad = with_shadow(im, slide.light) if shadow else (im, 0)
            canvas.alpha_composite(lay, (x - pad, y - pad)) if (x - pad >= 0 and y - pad >= 0) else canvas.paste(lay, (x - pad, y - pad), lay)
            base = os.path.splitext(os.path.basename(label))[0]
            p = os.path.join(tmp, f'{os.path.basename(out_png)[:-4]}-{n:02d}-{base}.png'); lay.save(p)
            pic = shapes.add_picture(p, Emu((x - pad) * EMU), Emu((y - pad) * EMU), Emu(lay.width * EMU), Emu(lay.height * EMU)); pic.name = base
        else:
            _, txt, x, y, w, size, font, color, lh, align, label = it
            f = ImageFont.truetype(FONTMAP[font], size); d = ImageDraw.Draw(canvas)
            lines = wrap(txt, f, w); asc, desc = f.getmetrics()
            for i, ln in enumerate(lines):
                lx = x if align == 'left' else (x + (w - f.getlength(ln)) / 2 if align == 'center' else x + w - f.getlength(ln))
                d.text((lx, y + i * size * lh), ln, font=f, fill=color)
            hgt = int(len(lines) * size * lh + desc)
            tb = shapes.add_textbox(Emu(x * EMU), Emu(y * EMU), Emu(w * EMU), Emu(hgt * EMU)); tb.name = label
            tf = tb.text_frame; tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
            fam, ital, bold = PPTFONT[font]
            for i, ln in enumerate(lines):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
                para.line_spacing = Pt(size * lh * 0.75)
                r = para.add_run(); r.text = ln; r.font.name = fam; r.font.size = Pt(size * 0.75)
                r.font.italic = ital; r.font.bold = bold; r.font.color.rgb = RGBColor(*color)
    if slide.notes: pptx_slide.notes_slide.notes_text_frame.text = slide.notes
    canvas.convert('RGB').save(out_png, quality=95)

def build(slides, w, h, pptx_path, prefix):
    prs = Presentation(); prs.slide_width = Emu(w * EMU); prs.slide_height = Emu(h * EMU)
    pngs = []
    for i, s in enumerate(slides, 1):
        ps = prs.slides.add_slide(prs.slide_layouts[6])
        out = os.path.join(HERE, f'{prefix}-{i}.png'); render(s, out, ps, prs); pngs.append(out)
    prs.save(pptx_path); return pngs

def contact(pngs, path, cols):
    ims = [Image.open(p) for p in pngs]; tw = 360; th = round(ims[0].height * tw / ims[0].width)
    rows = math.ceil(len(ims) / cols); g = 24
    c = Image.new('RGB', (cols * tw + (cols + 1) * g, rows * th + (rows + 1) * g), (28, 28, 28))
    for i, im in enumerate(ims):
        c.paste(im.convert('RGB').resize((tw, th), Image.LANCZOS), (g + (i % cols) * (tw + g), g + (i // cols) * (th + g)))
    c.save(path)

# ------------------------------------------------------------------ carousel
W, H = 1080, 1350; BY = 1205  # burrow height where it crosses slide edges
LOGO = ('s1-02-logo.png', 853, 5, 214)
def head(s, txt, small=None, size=104, w=900, y=150, small_y=None):
    s.text(txt, 72, y, w, size, label='headline')
    if small:
        f = ImageFont.truetype(HEAD, size); n = len(wrap(txt, f, w))
        s.text(small, 76, small_y or int(y + n * size * 1.06 + 64), 860, 34, 'small', SAGE, 1.3, label='small line')
def logo(s): s.cut(LOGO[0], LOGO[1], LOGO[2], LOGO[3], shadow=False, label='logo')
def burrow(s, pts): r, x, y = ribbon_layer(pts, W, H); s.cut(r, x, y, label='burrow')

car = []
# 1 midden, burrow starts below it and leaves right edge
s = Slide(W, H, 'Source: Charles Darwin, The Formation of Vegetable Mould through the Action of Worms, 1881. Darwin Online: http://darwin-online.org.uk/')
s.light = (10, 12); s.bg('s1-01-background.jpg')
head(s, "Worm holes are real. They're in your garden.", "Look for leaves standing up in the lawn. That's a worm's front door.", size=108)
burrow(s, [(470, 1010), (480, 1080), (560, 1160), (760, BY), (1100, BY)])
s.cut('s8-04-dark-soil.png', 300, 905, 360, label='soil mound')
s.cut('s4-04-moss.png', 40, 1010, 290, label='moss')
s.cut('midden.png', 330, 560, h=430, label='midden')
logo(s); car.append(s)
# 2 tall burrow cross-section with ruler
s = Slide(W, H, '')
s.light = (10, 12); s.bg('s1-01-background.jpg')
head(s, 'Nightcrawlers dig straight down, up to 2 meters.', size=96, w=480)
burrow(s, [(-20, BY), (380, BY), (660, 1150)])
s.cut('burrow-cross-section.png', 600, 380, h=930, label='burrow cross-section')
rul = Image.new('RGBA', (60, 930), (0, 0, 0, 0)); d = ImageDraw.Draw(rul)
d.rectangle([28, 0, 32, 929], fill=CREAM + (255,))
for i in range(0, 11): yy = int(i * 92.8); d.rectangle([12 if i % 5 == 0 else 20, yy, 48 if i % 5 == 0 else 40, yy + 3], fill=CREAM + (255,))
s.cut(rul, 935, 380, shadow=False, label='ruler')
s.text('2 m', 990, 1255, 80, 30, 'small', CREAM, label='ruler label')
logo(s); car.append(s)
# 3 microbe city: burrow crosses, magnifier on the wall
s = Slide(W, H, 'Source: Tiunov and Scheu 1999, burrow walls of Lumbricus terrestris: https://www.nies.go.jp/chiiki1/protoz/refere/id6999/6910.htm')
s.light = (-10, 12); s.bg('s1-01-background.jpg')
head(s, 'The burrow wall is a microbe city.', 'Up to 9 times more microbial activity than the soil around it.', size=104)
burrow(s, [(-20, BY), (1100, BY)])
lens = Image.new('RGBA', (40, 140), (0, 0, 0, 0)); ImageDraw.Draw(lens).line([(20, 0), (20, 139)], fill=CREAM + (255,), width=6)
s.cut(lens, 520, 1060, shadow=False, label='magnifier stem')
circ = Image.open(src('drilosphere-circle.png')).convert('RGBA').resize((560, 536), Image.LANCZOS)
ring = Image.new('RGBA', (600, 576), (0, 0, 0, 0)); ImageDraw.Draw(ring).ellipse([2, 2, 597, 573], fill=CREAM + (255,))
ring.alpha_composite(circ, (20, 20)); s.cut(ring, 240, 500, label='drilosphere magnifier')
s.cut('s8-05-earthworm.png', 790, 1090, 210, label='earthworm')
logo(s); car.append(s)
# 4 plumbing: rain down a vertical shaft into the burrow
s = Slide(W, H, 'Source: Capowiez et al. 2015, Biology and Fertility of Soils: https://doi.org/10.1007/s00374-015-1036-x')
s.light = (10, 12); s.bg('s1-01-background.jpg')
head(s, 'The burrows are plumbing.', 'Rain runs straight down the tunnels, deep into the soil.', size=110)
burrow(s, [[(-20, BY), (1100, BY)], [(560, 800), (560, BY)]])
for i, (rx, top, ln) in enumerate([(500, 560, 190), (535, 600, 260), (560, 545, 390), (585, 610, 250), (620, 570, 200)]):
    rl = Image.new('RGBA', (12, ln), (0, 0, 0, 0)); ImageDraw.Draw(rl).rounded_rectangle([3, 0, 8, ln - 1], 3, fill=CREAM + (230,))
    s.cut(rl, rx, top, shadow=False, label=f'rain {i + 1}')
s.cut('s8-03-wood-chips.png', 610, 735, 300, label='wood chips')
s.cut('s4-04-moss.png', 240, 745, 280, label='moss')
logo(s); car.append(s)
# 5 roots follow the burrow
s = Slide(W, H, 'Source: Kautz, Athmann and Köpke 2014: https://orgprints.org/23806/')
s.light = (-10, 12); s.bg('s1-01-background.jpg')
head(s, 'Roots follow the worms down.', 'They grow along the nitrogen-rich burrow walls into the subsoil.', size=110)
burrow(s, [(-20, BY), (520, BY)])
s.cut('root-in-burrow.png', 560, 520, h=800, label='root in burrow')
s.cut('s4-05-fern-fiddlehead.png', 780, 1040, 230, label='fiddlehead')
s.cut('s7-02-earthworm.png', 80, 960, 240, label='earthworm')
logo(s); car.append(s)
# 6 plowing cuts the pipes
s = Slide(W, H, 'Source: Briones and Schmidt 2017, Global Change Biology: https://doi.org/10.1111/gcb.13744')
s.light = (10, 12); s.bg('s1-01-background.jpg')
head(s, 'Plowing cuts the pipes.', 'No-till fields have 137% more earthworms than plowed ones.', size=110)
burrow(s, [(-20, BY), (560, BY)])
s.text('137%', 60, 640, 600, 250, 'num', CREAM, 1.0, label='137%')
s.cut('spade-cut.png', 600, 520, h=790, label='spade cut')
logo(s); car.append(s)

pngs = build(car, W, H, os.path.join(HERE, '21-10-2026-wed-ig-worm-holes.pptx'), 'worm-holes-slide')
contact(pngs, os.path.join(HERE, 'worm-holes-contact-sheet.png'), 3)

# ------------------------------------------------------------------ story
W, H = 1080, 1920
st = []
s = Slide(W, H); s.light = (10, 12); s.bg('s1-01-background.jpg')
s.text('Is World Earthworm Day even a thing?', 72, 300, 900, 128, label='headline')
s.cut('s8-04-dark-soil.png', 300, 1180, 520, label='dark soil')
s.cut('s4-04-moss.png', 40, 1300, 360, label='moss')
s.cut('s8-06-dry-leaves.png', 720, 1150, 330, label='dry leaves')
s.cut('s4-05-fern-fiddlehead.png', 90, 1050, 260, label='fiddlehead')
s.cut('s7-02-earthworm.png', 420, 780, 400, label='earthworm')
s.cut('s8-05-earthworm.png', 660, 1440, 320, label='earthworm 2')
s.cut('s1-02-logo.png', 853, 160, 214, shadow=False, label='logo'); st.append(s)

s = Slide(W, H, 'Earthworm Society of Britain, World Earthworm Day: https://www.earthwormsoc.org.uk/worldwormday\nDarwin Correspondence Project, letter 13460 (3,500 Worms, 600 Origin): https://www.darwinproject.ac.uk/view/letters/DCP-LETT-13460')
s.light = (10, 12); s.bg('s1-01-background.jpg')
s.text('Yes! It honors Darwin. His worm book came out in October 1881.', 72, 270, 920, 96, label='headline')
s.text("At his publisher's sale that November, it sold 3,500 copies. Origin of Species sold 600.", 76, 700, 880, 36, 'small', SAGE, 1.3, label='small line')
card = Image.new('RGBA', (380, 560), CREAM + (255,))
tc = Image.open(src('darwin-fig3-tower-casting.png')).convert('RGBA'); tc.thumbnail((300, 470))
card.alpha_composite(tc, ((380 - tc.width) // 2, (560 - tc.height) // 2))
s.cut(card, 620, 880, label='Darwin woodcut, fig 3')
s.cut('piano-pots.png', 70, 880, 470, label='piano pots')
s.cut('paper-triangles.png', 60, 1390, 420, label='paper triangles')
s.cut('worm-stone.png', 540, 1480, 470, label='worm stone')
s.cut('s1-02-logo.png', 853, 140, 214, shadow=False, label='logo'); st.append(s)

s = Slide(W, H); s.light = (10, 12); s.bg('s1-01-background.jpg')
s.text('Which earthworm are you today?', 72, 290, 940, 112, label='headline')
cols = [('worm-epigeic.png', 'Epigeic.', 'Lives in the leaf litter and eats fast.'),
        ('worm-endogeic.png', 'Endogeic.', 'Stays in the topsoil and eats its way through it.'),
        ('worm-anecic.png', 'Anecic.', 'Digs deep vertical burrows and comes up at night.')]
cw = 312; gx = 36
for i, (f, name, desc) in enumerate(cols):
    x0 = gx + i * (cw + 18)
    im = Image.open(src(f)).convert('RGBA'); sc = min(300 / im.width, 380 / im.height); iw, ih = int(im.width * sc), int(im.height * sc)
    s.cut(f, x0 + (cw - iw) // 2, 970 - ih, iw, ih, label=name)
    s.text(name, x0, 995, cw, 60, 'head', CREAM, 1.0, 'center', label=name + ' label')
    s.text(desc, x0 + 6, 1075, cw - 12, 31, 'small', SAGE, 1.3, 'center', label=name + ' text')
s.cut('s1-02-logo.png', 853, 140, 214, shadow=False, label='logo')
s.cut('s8-06-dry-leaves.png', 60, 1600, 260, label='dry leaves'); s.cut('s4-04-moss.png', 720, 1620, 300, label='moss'); st.append(s)

pngs = build(st, W, H, os.path.join(HERE, '21-10-2026-wed-ig-story-earthworm.pptx'), 'story-frame')
contact(pngs, os.path.join(HERE, 'story-contact-sheet.png'), 3)
print('done')
