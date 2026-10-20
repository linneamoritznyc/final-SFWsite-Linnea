"""World Earthworm Day Instagram Story, 3 frames, 1080 x 1920. Silly: every worm wears a hat.
Words exactly as in the October doc. Worm: social/cutouts-soil/earthworm.png; hats are drawn here.
Keep text inside y 250..1670 (Story safe zone); frame 3 leaves room for the poll sticker.

  python3 social/earthworm-day-2026-10-21/story/build_story.py -> frame-1..3.png, story-contact.png, earthworm-story-editable.pptx
"""
import os, json, math, subprocess
import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageEnhance

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
W, H = 1080, 1920
DEEP, GREEN, CREAM, GOLD, SOIL = "#22371F", "#156826", "#F4F1EA", "#C9A227", "#4F3433"
LAY = os.path.join(HERE, "layers"); os.makedirs(LAY, exist_ok=True)
for f in os.listdir(LAY): os.remove(os.path.join(LAY, f))


def rgb(h): return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))


# ---------------- worms and hats ----------------
WORM = Image.open(os.path.join(ROOT, "social", "cutouts-soil", "earthworm.png")).convert("RGBA").rotate(180, expand=True)
TIP = (8, 250)   # the head end after turning the worm over: pointing up, on the left


def tint(im, sat=1.0, bright=1.0, warm=(0, 0, 0)):
    a = im.getchannel("A"); c = ImageEnhance.Color(im.convert("RGB")).enhance(sat); c = ImageEnhance.Brightness(c).enhance(bright)
    arr = np.clip(np.asarray(c, np.float32) + np.array(warm, np.float32), 0, 255).astype(np.uint8)
    out = Image.fromarray(arr).convert("RGBA"); out.putalpha(a); return out


def party_hat(s):
    S = 4; w, h = int(70 * s), int(110 * s); im = Image.new("RGBA", (w * S, (h + 30) * S), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    cone = [(w * S // 2, 22 * S), (0, (h + 22) * S), (w * S, (h + 22) * S)]
    d.polygon(cone, fill=rgb(GREEN) + (255,))
    m = Image.new("L", im.size, 0); ImageDraw.Draw(m).polygon(cone, fill=255)
    stripes = Image.new("RGBA", im.size, (0, 0, 0, 0)); sd = ImageDraw.Draw(stripes)
    for k in range(-6, 8):
        y0 = k * 30 * S; sd.polygon([(0, y0 + 60 * S), (w * S, y0), (w * S, y0 + 13 * S), (0, y0 + 73 * S)], fill=rgb(GOLD) + (255,))
    stripes.putalpha(Image.fromarray(np.minimum(np.asarray(stripes.getchannel("A")), np.asarray(m)))); im.alpha_composite(stripes)
    d.ellipse((w * S // 2 - 16 * S, 0, w * S // 2 + 16 * S, 32 * S), fill=(232, 96, 76, 255))
    return im.resize((w, h + 30), Image.LANCZOS)


def straw_hat(s):
    S = 4; w, h = int(150 * s), int(80 * s); im = Image.new("RGBA", (w * S, h * S), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    straw, dark = (232, 196, 120, 255), (196, 152, 82, 255)
    d.ellipse((0, h * S * .5, w * S, h * S * .95), fill=straw)
    for k in range(1, 6): d.ellipse((w * S * .05 * k, h * S * (.5 + .045 * k), w * S * (1 - .05 * k), h * S * (.95 - .045 * k)), outline=dark, width=3 * S)
    d.rounded_rectangle((w * S * .28, h * S * .08, w * S * .72, h * S * .74), radius=26 * S, fill=straw)
    for y in range(int(h * S * .16), int(h * S * .7), 9 * S): d.line((w * S * .3, y, w * S * .7, y), fill=dark, width=2 * S)
    d.rectangle((w * S * .28, h * S * .52, w * S * .72, h * S * .64), fill=(196, 72, 62, 255))
    return im.resize((w, h), Image.LANCZOS)


def miner_helmet(s):
    S = 4; w, h = int(130 * s), int(100 * s); im = Image.new("RGBA", (w * S, h * S), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    d.pieslice((w * S * .1, h * S * .18, w * S * .9, h * S * 1.3), 180, 360, fill=(250, 196, 40, 255))
    d.rectangle((0, h * S * .7, w * S, h * S * .8), fill=(226, 170, 24, 255))
    d.line((w * S * .5, h * S * .2, w * S * .5, h * S * .72), fill=(226, 170, 24, 255), width=5 * S)
    d.rounded_rectangle((w * S * .38, h * S * .32, w * S * .62, h * S * .58), radius=10 * S, fill=(70, 70, 70, 255))
    d.ellipse((w * S * .42, h * S * .35, w * S * .58, h * S * .55), fill=(255, 246, 200, 255))
    return im.resize((w, h), Image.LANCZOS)


def lamp_beam(length=420, spread=70):
    im = Image.new("RGBA", (length, spread * 2 + 40), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    d.polygon([(0, spread + 20), (length, 0), (length, spread * 2 + 40)], fill=(255, 238, 170, 110))
    return im.filter(ImageFilter.GaussianBlur(10))


def worm_with_hat(kind, scale):
    w = WORM.resize((round(WORM.width * scale), round(WORM.height * scale)), Image.LANCZOS)
    if kind == "epigeic": w = tint(w, 1.35, 1.0, (26, -6, -8)); hat = party_hat(scale * 1.0); rot = 10
    elif kind == "endogeic": w = tint(w, .45, 1.22, (18, 6, 10)); hat = straw_hat(scale * 1.0); rot = 6
    else: w = tint(w, 1.05, .72, (-4, -8, -6)); hat = miner_helmet(scale * 1.0); rot = 4
    hat = hat.rotate(rot, expand=True, resample=Image.BICUBIC)
    pad = hat.height + 20
    out = Image.new("RGBA", (w.width + hat.width, w.height + pad), (0, 0, 0, 0))
    ox = hat.width // 2; out.alpha_composite(w, (ox, pad))
    tx, ty = ox + TIP[0] * scale, pad + TIP[1] * scale
    out.alpha_composite(hat, (round(tx - hat.width / 2 + 14 * scale), round(ty - hat.height + 30 * scale)))
    return out.crop(out.getchannel("A").getbbox())


# ---------------- backgrounds ----------------
def grad(stops, h=H):
    y = np.linspace(0, 1, h); pos = [p for p, _ in stops]; cols = np.array([rgb(c) for _, c in stops], np.float32)
    g = np.stack([np.interp(y, pos, cols[:, k]) for k in range(3)], -1)
    a = np.repeat(g[:, None, :], W, 1) + np.random.default_rng(1).normal(0, 2.5, (h, W, 1))
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))


def cover(path, w, h, fx=.5, fy=.5):
    im = Image.open(path).convert("RGB"); s = max(w / im.width, h / im.height)
    im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    x, y = round((im.width - w) * fx), round((im.height - h) * fy); return im.crop((x, y, x + w, y + h))


def soil(w, h, dark=0.0, seed=3):
    src = Image.open(os.path.join(ROOT, "img", "new-2026-10", "R5A_4027.jpg")).convert("RGB").crop((60, 500, 1499, 1000))
    a = np.asarray(cover_img(src, w, h), np.float32) * (1 - dark)
    im = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8)).convert("RGBA")
    rnd = np.random.default_rng(seed); edge = Image.new("L", (w, h), 255); d = ImageDraw.Draw(edge)
    d.polygon([(0, 0)] + [(x, float(rnd.uniform(0, 16))) for x in range(0, w + 20, 20)] + [(w, 0)], fill=0)
    im.putalpha(edge.filter(ImageFilter.GaussianBlur(1))); return im


def cover_img(im, w, h):
    s = max(w / im.width, h / im.height); im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    x, y = (im.width - w) // 2, (im.height - h) // 2; return im.crop((x, y, x + w, y + h))


def shadow(im, off=(8, 12), blur=12, a=.3):
    pad = blur * 3 + 14; out = Image.new("RGBA", (im.width + 2 * pad, im.height + 2 * pad), (0, 0, 0, 0))
    sh = Image.new("RGBA", im.size, (40, 28, 20, 255)); sh.putalpha(im.getchannel("A").point(lambda v: int(v * a)))
    l = Image.new("RGBA", out.size, (0, 0, 0, 0)); l.alpha_composite(sh, (pad + off[0], pad + off[1]))
    out.alpha_composite(l.filter(ImageFilter.GaussianBlur(blur))); out.alpha_composite(im, (pad, pad)); return out, pad


# ---------------- frames ----------------
FRAMES = []


class Frame:
    def __init__(self, n):
        self.n, self.items = n, []; FRAMES.append(self)

    def img(self, name, im, x, y, jpg=False, center=False, shade=False):
        if shade: im, pad = shadow(im); x -= pad; y -= pad
        if center: x, y = round(x - im.width / 2), round(y - im.height / 2)
        box = (max(0, -x), max(0, -y), min(im.width, W - x), min(im.height, H - y)); im = im.crop(box); x += box[0]; y += box[1]
        fn = f"f{self.n}-{len(self.items) + 1:02d}-{name}." + ("jpg" if jpg else "png")
        (im.convert("RGB").save(os.path.join(LAY, fn), quality=90) if jpg else im.save(os.path.join(LAY, fn), optimize=True))
        self.items.append(dict(type="image", name=name, file="layers/" + fn, x=x, y=y, w=im.width, h=im.height))

    def text(self, name, t, x, y, w, h, size, color, font="Montserrat", weight=700, lh=1.08, align="left"):
        self.items.append(dict(type="text", name=name, text=t, x=x, y=y, w=w, h=h, size=size, color=color, font=font, weight=weight, lh=lh, align=align))


logo = Image.open(os.path.join(ROOT, "img", "sfwlogo-240.png")).convert("RGBA"); logo = logo.resize((150, round(150 * logo.height / logo.width)), Image.LANCZOS)
SKY = [(0, "#FFD27F"), (.55, "#FFE9C2"), (1, "#FFF6E6")]

# Frame 1: the question, three hatted worms popping up out of the soil
f = Frame(1)
f.img("background", grad(SKY), 0, 0, jpg=True)
f.text("headline", "Is World Earthworm Day even a thing?", 84, 330, 912, 500, 110, DEEP, lh=1.04)
f.img("soil", soil(W, 640), 0, H - 640)
for k, (kind, sc, cx) in enumerate([("epigeic", .62, 230), ("endogeic", .78, 560), ("anecic", 1.0, 880)]):
    wm = worm_with_hat(kind, sc); f.img(f"worm-{kind}", wm, round(cx - wm.width / 2), H - 640 - wm.height + 200, shade=True)
f.img("logo", logo, W - 84 - 150, 1700 - logo.height)

# Frame 2: yes, it's real
f = Frame(2)
f.img("background", grad([(0, "#FFF6E6"), (1, "#FFE3B3")]), 0, 0, jpg=True)
f.text("yes", "Yes, it's real.", 84, 300, 912, 140, 104, GREEN, lh=1.0)
f.text("story", "In 2016 the Earthworm Society of Britain nominated 21 October as World Earthworm Day, and they chose October to honor Charles Darwin, since his book on worms was published that month.",
       84, 470, 912, 560, 50, SOIL, font="Source Sans 3", weight=600, lh=1.3)
f.img("soil", soil(W, 520, .05, 7), 0, H - 520)
wm = worm_with_hat("anecic", 1.15); f.img("worm-anecic", wm, 640, H - 520 - wm.height + 260, shade=True)
beam = lamp_beam(); f.img("lamp-beam", beam.transpose(Image.FLIP_LEFT_RIGHT).rotate(-12, expand=True), 250, H - 520 - wm.height + 140)

# Frame 3: which earthworm are you today? Three habitat panels, poll sticker space under them
f = Frame(3)
f.img("background", grad(SKY), 0, 0, jpg=True)
f.text("headline", "Which earthworm are you today?", 84, 250, 912, 260, 92, DEEP, lh=1.04)
PANELS = [("epigeic", "Epigeic. Lives in the leaf litter and eats fast.", "leaf"),
          ("endogeic", "Endogeic. Stays in the topsoil and eats its way through it.", "top"),
          ("anecic", "Anecic. Digs deep vertical burrows and comes up at night.", "deep")]
top0, ph, gap = 500, 300, 28
for k, (kind, label, hab) in enumerate(PANELS):
    y = top0 + k * (ph + gap); pw = W - 2 * 84
    if hab == "leaf":
        base = Image.new("RGB", (pw, ph), (214, 170, 110))
        leaves = Image.open(os.path.join(ROOT, "social", "cutouts-soil", "dry-leaves-twigs.png")).convert("RGBA")
        for j in range(4):
            l = leaves.resize((260, round(260 * leaves.height / leaves.width))).rotate(j * 70, expand=True)
            base.paste(l, (j * 230 - 40, -20 + (j % 2) * 60), l)
    elif hab == "top": base = soil(pw, ph, -.12, 11).convert("RGB")
    else: base = soil(pw, ph, .45, 13).convert("RGB")
    m = Image.new("L", (pw, ph), 0); ImageDraw.Draw(m).rounded_rectangle((0, 0, pw - 1, ph - 1), radius=36, fill=255)
    panel = base.convert("RGBA"); panel.putalpha(m)
    if hab == "deep":  # a dark burrow running down
        bd = ImageDraw.Draw(panel); bd.line([(pw - 250, 0), (pw - 240, ph)], fill=(28, 18, 12, 230), width=46)
    shade_ = np.zeros((ph, pw, 4), np.uint8); shade_[..., 3] = (np.linspace(0, 1, pw)[None, :] ** 2 * 0 + np.linspace(1, 0, pw)[None, :] * 150).astype(np.uint8)
    sh = Image.fromarray(shade_); sh.putalpha(Image.fromarray(np.minimum(np.asarray(sh.getchannel("A")), np.asarray(m))))
    panel.alpha_composite(sh)                                                  # darker on the left so the label reads
    f.img(f"panel-{kind}", panel, 84, y, shade=True)
    f.text(f"label-{kind}", label, 120, y + 40, 520, 230, 40, CREAM, font="Source Sans 3", weight=600, lh=1.22)
    wm = worm_with_hat(kind, {"epigeic": .5, "endogeic": .56, "anecic": .66}[kind])
    f.img(f"worm-{kind}", wm, W - 84 - 40 - wm.width, y + ph - wm.height - 14, shade=True)

# ---------------- output: HTML render, PNGs, PPTX ----------------
fonts = "file://" + os.path.join(ROOT, "fonts")
css = f"""@font-face{{font-family:"Montserrat";src:url("{fonts}/montserrat-latin-variable.woff2");font-weight:100 900}}
@font-face{{font-family:"Source Sans 3";src:url("{fonts}/source-sans-3-latin-600-normal.woff2");font-weight:600}}
*{{margin:0}}body{{background:#555;display:grid;gap:40px;padding:40px}}.slide{{position:relative;width:{W}px;height:{H}px;overflow:hidden}}
.slide img{{position:absolute;display:block}}.slide p{{position:absolute}}"""
out = [f"<!doctype html><html><head><meta charset='utf-8'><title>Earthworm story</title><style>{css}</style></head><body>"]
for fr in FRAMES:
    out.append("<section class='slide'>")
    for it in fr.items:
        if it["type"] == "image": out.append(f"<img src='{it['file']}' style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;height:{it['h']}px' alt=''>")
    for it in fr.items:
        if it["type"] == "text":
            out.append(f"<p style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;font-family:\"{it['font']}\";font-size:{it['size']}px;"
                       f"font-weight:{it['weight']};line-height:{it['lh']};color:{it['color']}'>{it['text']}</p>")
    out.append("</section>")
open(os.path.join(HERE, "slides.html"), "w").write("\n".join(out + ["</body></html>"]))
json.dump([dict(n=fr.n, items=fr.items) for fr in FRAMES], open(os.path.join(HERE, "layout.json"), "w"), indent=1)
js = os.path.join(HERE, "_render.cjs")
open(js, "w").write("""const { chromium } = require("playwright"); const { pathToFileURL } = require("node:url"); const path = require("node:path");
(async () => { const b = await chromium.launch(); const p = await b.newPage({ viewport: { width: 1200, height: 2100 } });
await p.goto(pathToFileURL(path.join(__dirname, "slides.html")).href); await p.evaluate(() => document.fonts.ready);
const s = await p.$$(".slide"); for (const [i, e] of s.entries()) await e.screenshot({ path: path.join(__dirname, `frame-${i + 1}.png`) });
await b.close(); })();""")
subprocess.run(["node", js], check=True); os.remove(js)
sheet = Image.new("RGB", (3 * 400 + 80, 711 + 40), (31, 31, 31))
for i in range(3): sheet.paste(Image.open(os.path.join(HERE, f"frame-{i + 1}.png")).convert("RGB").resize((400, 711), Image.LANCZOS), (20 + i * 420, 20))
sheet.save(os.path.join(HERE, "story-contact.png"))

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
E = 914400 / 108
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(round(W * E)), Emu(round(H * E))
for fr in FRAMES:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for it in fr.items:
        if it["type"] == "image":
            sl.shapes.add_picture(os.path.join(HERE, it["file"]), *[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]).name = it["name"]
    for it in fr.items:
        if it["type"] != "text": continue
        tb = sl.shapes.add_textbox(*[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]); tb.name = it["name"]
        tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.line_spacing = it["lh"]; r = p.add_run(); r.text = it["text"]; fo = r.font
        fo.name = "Source Sans Pro" if it["font"] == "Source Sans 3" else it["font"]; fo.size = Pt(round(it["size"] * 72 / 108, 1))
        fo.bold = it["weight"] >= 600; fo.color.rgb = RGBColor(*rgb(it["color"]))
prs.save(os.path.join(HERE, "earthworm-story-editable.pptx")); print("done")
