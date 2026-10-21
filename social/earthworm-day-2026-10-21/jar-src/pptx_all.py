import json, re, os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
HERE = os.path.dirname(os.path.abspath(__file__)); os.chdir(HERE)
OUT = '/home/user/final-sfwsite-linnea/social/earthworm-day-2026-10-21/21-10-2026-wed-ig-earthworm-jar.pptx'
E = lambda px: Emu(int(round(px * 9525)))
INK = RGBColor(0x4F, 0x34, 0x33)
prs = Presentation(); prs.slide_width, prs.slide_height = E(1080), E(1350)
for spec in json.load(open('layers/specs.json')):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(255, 255, 255)
    for name, path in spec['pics']:
        im = Image.open(path); bb = im.getbbox(); c = im.crop(bb); p = path.replace('.png', '-crop.png'); c.save(p)
        sl.shapes.add_picture(p, E(bb[0]), E(bb[1]), E(c.width), E(c.height)).name = name
    for n, (p0, p1, bend) in enumerate(spec['arrows']):
        mx, my = (p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2
        dx, dy = p1[0] - p0[0], p1[1] - p0[1]
        cx, cy = mx - dy * bend, my + dx * bend
        pts = [((1-t)**2*p0[0] + 2*(1-t)*t*cx + t*t*p1[0], (1-t)**2*p0[1] + 2*(1-t)*t*cy + t*t*p1[1]) for t in [k/24 for k in range(25)]]
        fb = sl.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
        sh = fb.convert_to_shape(); sh.name = f'Arrow {n+1}'
        sh.fill.background(); sh.line.color.rgb = INK; sh.line.width = Pt(2)
        ln = sh.line._get_or_add_ln(); etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    for t in spec['texts']:
        t.setdefault('lh', 1.25); t.setdefault('weight', 500); t.setdefault('family', 'body'); t.setdefault('color', '#5b4c46'); t.setdefault('align', 'left')
        lines = re.split(r'<br\s*/?>', t['text'])
        tb = sl.shapes.add_textbox(E(t['x']), E(t['y'] - 4), E(t['w'] + 30), E(t['size'] * t['lh'] * len(lines) + 10))
        tb.name = re.sub('<[^>]+>', ' ', t['text'])[:40]
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fam, bold = (('Fredoka', t['weight'] >= 600) if t['family'] == 'title' else ('Montserrat', t['weight'] >= 600))
        al = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[t['align']]
        for i, l in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = t['lh']; p.alignment = al
            r = p.add_run(); r.text = l.replace('&nbsp;', ' ')
            r.font.size = Pt(t['size'] * 0.75); r.font.name = fam; r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(t['color'].lstrip('#'))
    lg = Image.open('logo.png')
    sl.shapes.add_picture('logo.png', E(890), E(1200), E(150), E(150 * lg.height / lg.width)).name = 'SFW logo'
    sl.notes_slide.notes_text_frame.text = 'Fonts: Fredoka (titles, labels) and Montserrat (body), free on Google Fonts. Illustration drawn in code (no AI). Arrows are editable lines.'
prs.save(OUT); print(OUT, len(prs.slides.__iter__.__self__._sldIdLst))
