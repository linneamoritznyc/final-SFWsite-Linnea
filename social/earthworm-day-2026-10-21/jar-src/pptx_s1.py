import json, re
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
OUT = '/home/user/final-sfwsite-linnea/social/earthworm-day-2026-10-21/21-10-2026-wed-ig-earthworm-jar.pptx'
E = lambda px: Emu(int(round(px * 9525)))
prs = Presentation(); prs.slide_width, prs.slide_height = E(1080), E(1350)
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(255, 255, 255)
spec = json.load(open('layers/s1.json'))

def pic(path, name):
    im = Image.open(path); bb = im.getbbox(); c = im.crop(bb); p = f'layers/{name}-crop.png'; c.save(p)
    s = sl.shapes.add_picture(p, E(bb[0]), E(bb[1]), E(c.width), E(c.height)); s.name = name

pic('layers/jar.png', 'Jar')
pic('layers/casts.png', 'Worm casts')
for i in (1, 2, 3):
    pic(f'layers/worm{i}.png', f'Worm {i}')

INK = RGBColor(0x4F, 0x34, 0x33)
for n, (p0, p1, bend) in enumerate(spec['arrows']):
    mx, my = (p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2
    dx, dy = p1[0] - p0[0], p1[1] - p0[1]
    cx, cy = mx - dy * bend, my + dx * bend
    pts = [((1-t)**2*p0[0] + 2*(1-t)*t*cx + t*t*p1[0], (1-t)**2*p0[1] + 2*(1-t)*t*cy + t*t*p1[1]) for t in [k/24 for k in range(25)]]
    fb = sl.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
    sh = fb.convert_to_shape(); sh.name = f'Arrow {n+1}'
    sh.fill.background(); sh.line.color.rgb = INK; sh.line.width = Pt(2)
    ln = sh.line._get_or_add_ln()
    ln.append(etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med'))
    sh.shadow.inherit = False

def font_for(t):
    if t['family'] == 'title':
        return ('Fredoka', True) if t['weight'] >= 600 else ('Fredoka', False)
    return ('Montserrat', False)

for t in spec['texts']:
    t.setdefault('lh', 1.25); t.setdefault('weight', 500); t.setdefault('family', 'body'); t.setdefault('color', '#5b4c46')
    lines = re.split(r'<br\s*/?>', t['text'])
    nl = len(lines)
    h = t['size'] * t['lh'] * nl + 10
    tb = sl.shapes.add_textbox(E(t['x']), E(t['y'] - 4), E(t['w'] + 40), E(h))
    tb.name = re.sub('<[^>]+>', ' ', t['text'])[:40]
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    fam, bold = font_for(t)
    col = RGBColor.from_string(t['color'].lstrip('#'))
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = t['lh']; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = l
        r.font.size = Pt(t['size'] * 0.75); r.font.name = fam; r.font.bold = bold; r.font.color.rgb = col

lg = Image.open('logo.png'); w = 150; h = w * lg.height / lg.width
s = sl.shapes.add_picture('logo.png', E(890), E(1200), E(w), E(h)); s.name = 'SFW logo'
sl.notes_slide.notes_text_frame.text = 'Fonts: Fredoka (titles, labels) and Montserrat (subtitle), both free on Google Fonts. Every layer is its own object: jar, each worm, casts, arrows (editable lines), text, logo.'
prs.save(OUT); print(OUT)
