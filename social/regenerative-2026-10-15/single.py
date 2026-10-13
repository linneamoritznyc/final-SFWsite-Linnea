"""'What does Regenerative actually mean?' as ONE Instagram image (1080 x 1350), three design options.
Website typography: Montserrat Bold headings, Source Sans 3 small text (named Source Sans Pro in the
PowerPoint, the version Canva has).

  python3 social/regenerative-2026-10-15/single.py
    -> single-1/, single-2/, single-3/ (post.png, layers, editable .pptx) and single-compare.png
"""
import os
from PIL import Image, ImageDraw, ImageFont
import variants as v
from variants import Deck, W, H, DEEP, GREEN, LIGHT, CREAM, rgb, cut, ragged, soil_strip, logo_img, FLOW, NEW, HEAD, SMALL1, SMALL5, Q

SOIL = "#4F3433"
HERE = os.path.dirname(os.path.abspath(__file__))


class Single(Deck):
    def __init__(self, key):
        super().__init__(key); self.dir = os.path.join(HERE, f"single-{key}")
        self.lay = os.path.join(self.dir, "layers"); os.makedirs(self.lay, exist_ok=True)

    def write(self):
        super().write()
        os.replace(os.path.join(self.dir, "slide-1.png"), os.path.join(self.dir, "post.png"))
        for f in ("contact-sheet.png",):
            if os.path.exists(os.path.join(self.dir, f)): os.remove(os.path.join(self.dir, f))
        p = os.path.join(self.dir, f"regenerative-option-{self.key}-canva.pptx")
        os.replace(p, os.path.join(self.dir, f"regenerative-single-{self.key}-canva.pptx"))
        # Canva has Source Sans Pro, not Source Sans 3
        from pptx import Presentation
        prs = Presentation(os.path.join(self.dir, f"regenerative-single-{self.key}-canva.pptx"))
        for sh in prs.slides[0].shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.name == "Source Sans 3": r.font.name = "Source Sans Pro"
        prs.save(os.path.join(self.dir, f"regenerative-single-{self.key}-canva.pptx"))


def questions(d, s, x, y, w, size, before_c, after_c, gap=34, rule=None):
    """The three either/or questions, stacked; returns the y after the last one."""
    for i, (before, after) in enumerate(Q):
        per_line = w / (size * 0.56)
        lines = -(-(len(before) + len(after) + 1) // int(per_line))
        h = round(size * 1.12 * lines)
        d.text(s, f"question-{i + 1}", [(before + " ", before_c, False), (after, after_c, False)], x, y, w, h, size, lh=1.12)
        y += h + gap
        if rule and i < 2:
            d.img(s, f"rule-{i + 1}", Image.new("RGBA", (w, 3), rgb(rule) + (255,)), x, y - gap // 2 - 1)
    return y


# 1. Cream page on a torn strip of real soil, website colours
d = Single("1"); s = d.slide(CREAM)
d.img(s, "logo", logo_img(130), W - 72 - 130, 40)
d.text(s, "headline", [(HEAD, SOIL, False)], 72, 120, 936, 260, 76, lh=1.1)
d.text(s, "small-line", [(SMALL1, GREEN, False)], 74, 390, 800, 56, 44, font="Source Sans 3", weight=600)
y = questions(d, s, 72, 490, 936, 50, SOIL, GREEN, rule="#D9D0BF")
d.text(s, "cta", [(SMALL5, SOIL, False)], 74, y + 6, 560, 110, 38, font="Source Sans 3", lh=1.25)
d.img(s, "soil-strip", soil_strip(270, 7), 0, H - 270)
im = cut("fern-fiddlehead.png", w=200); d.img(s, "fern-fiddlehead", im, 940, H - 270 - im.height // 2 + 50, center=True)
im = cut("mushrooms.png", w=180); d.img(s, "mushrooms", im, 780, H - 270 - im.height // 2 + 50, center=True)
d.write()

# 2. Deep green ground with flowers from the four workshop regions
d = Single("2"); s = d.slide(DEEP)
d.img(s, "logo", logo_img(100, disc=True), W - 72 - 125, 40)
d.text(s, "headline", [(HEAD, CREAM, False)], 72, 120, 936, 260, 76, lh=1.08)
d.text(s, "small-line", [(SMALL1, LIGHT, False)], 74, 390, 800, 56, 44, font="Source Sans 3", weight=600)
y = questions(d, s, 72, 490, 936, 50, CREAM, LIGHT, rule="#3A5235")
d.text(s, "cta", [(SMALL5, CREAM, False)], 74, y + 6, 520, 110, 38, font="Source Sans 3", lh=1.25)
for nm, f, h_, cx, cy, r in [("poppy", "norfolk-common-poppy.png", 330, 760, 1150, -6), ("marigold", "tamil-nadu-marigold.png", 280, 610, 1220, 6),
                             ("paintbrush", "santa-fe-indian-paintbrush.png", 380, 900, 1120, -4), ("heliconia", "costa-rica-heliconia.png", 420, 1010, 1130, 8)]:
    d.img(s, "flower-" + nm, cut(f, h=h_, rot=r, folder=FLOW), cx, cy, center=True)
d.write()

# 3. One photo: the August 2023 high-five over the compost, type on cream below a torn edge
d = Single("3"); s = d.slide(CREAM)
ph = Image.open(os.path.join(NEW, "students-high-five-at-compost.jpg")).convert("RGB")
ph = v.cover(os.path.join(NEW, "students-high-five-at-compost.jpg"), W, 560, .35).convert("RGBA"); ph.putalpha(ragged(W, 560, "b", 91, 14))
d.img(s, "photo-high-five-at-compost", ph, 0, 0)
d.img(s, "logo", logo_img(110, disc=True), W - 72 - 138, 40)
d.text(s, "headline", [(HEAD, SOIL, False)], 72, 590, 936, 160, 64, lh=1.1)
d.text(s, "small-line", [(SMALL1, GREEN, False)], 74, 745, 800, 50, 40, font="Source Sans 3", weight=600)
questions(d, s, 72, 830, 936, 40, SOIL, GREEN, gap=24, rule="#D9D0BF")
d.text(s, "cta", [(SMALL5, SOIL, False)], 74, 1240, 936, 60, 34, font="Source Sans 3", lh=1.25)
d.write()

# side by side
tw, th, g = 480, 600, 30
cmp_ = Image.new("RGB", (g + 3 * (tw + g), th + 2 * g + 40), (31, 31, 31)); dr = ImageDraw.Draw(cmp_)
f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 26)
for i, k in enumerate("123"):
    x = g + i * (tw + g); dr.text((x, 18), f"Option {k}", font=f, fill=(240, 240, 240))
    cmp_.paste(Image.open(os.path.join(HERE, f"single-{k}", "post.png")).convert("RGB").resize((tw, th), Image.LANCZOS), (x, 60))
cmp_.save(os.path.join(HERE, "single-compare.png")); print("done")
