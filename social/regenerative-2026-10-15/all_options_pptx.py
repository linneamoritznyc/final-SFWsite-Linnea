"""Every Regenerative option in one editable PowerPoint, newest first. Each slide's notes say which option it is.
Run the builders first (sunny.py, rules.py, concepts.py, single.py, variants.py, build.py), then:
  python3 social/regenerative-2026-10-15/all_options_pptx.py -> regenerative-all-options-canva.pptx
"""
import os, json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
W, H, E = 1080, 1350, 914400 / 108
SETS = [("more-a", "More A: poppy field"), ("more-b", "More B: low sun"), ("more-c", "More C: flower frame"),
        ("more-d", "More D: golden sunrise"), ("sunny-1", "Sunny 1: poppy field, Wild Ken Hill"), ("sunny-2", "Sunny 2: cover crops, Capri"),
        ("sunny-3", "Sunny 3: morning sun, flowers"),
        ("rules-1", "Text-post rules 1: cracked to alive"), ("rules-3", "Text-post rules 3: soil in the letters"),
        ("concept-1", "Concept 1: cracked to alive"), ("concept-3", "Concept 3: soil in the letters"),
        ("single-1", "Single image 1: cream and soil"), ("single-2", "Single image 2: deep green, flowers"),
        ("single-3", "Single image 3: one photo"),
        (".", "Carousel A"), ("option-b", "Carousel B: deep green ground"), ("option-c", "Carousel C: the soil line"),
        ("option-d", "Carousel D: specimen sheet"), ("option-e", "Carousel E: one photo only")]
CANVA_FONT = {"Source Sans 3": "Source Sans Pro"}


def rgb(h): return RGBColor(*(int(h[i:i + 2], 16) for i in (1, 3, 5)))


prs = Presentation(); prs.slide_width, prs.slide_height = Emu(round(W * E)), Emu(round(H * E))
count = 0
for folder, label in SETS:
    base = os.path.join(HERE, folder)
    slides = json.load(open(os.path.join(base, "layout.json")))
    for s in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[6]); count += 1
        sl.notes_slide.notes_text_frame.text = f"{label}, slide {s['n']} of {len(slides)}" if len(slides) > 1 else label
        for it in s["items"]:
            if it["type"] == "image":
                sl.shapes.add_picture(os.path.join(base, it["file"]), *[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]).name = it["name"]
        for it in s["items"]:
            if it["type"] != "text": continue
            tb = sl.shapes.add_textbox(*[Emu(round(it[k] * E)) for k in ("x", "y", "w", "h")]); tb.name = it["name"]
            tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            for t, c, nl in it["runs"]:
                p.line_spacing = it["lh"]; p.alignment = PP_ALIGN.CENTER if it.get("align") == "center" else PP_ALIGN.LEFT
                r = p.add_run(); r.text = t; f = r.font
                f.name = CANVA_FONT.get(it["font"], it["font"]); f.size = Pt(round(it["size"] * 72 / 108, 1))
                f.bold = it["weight"] >= 600; f.color.rgb = rgb(c)
                if nl: p = tf.add_paragraph()
out = os.path.join(HERE, "regenerative-all-options-canva.pptx")
prs.save(out); print(count, "slides ->", out, round(os.path.getsize(out) / 1e6, 1), "MB")
