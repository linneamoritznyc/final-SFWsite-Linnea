"""'What does Regenerative actually mean?' 5-slide Instagram carousel, Thursday 15 October 2026.

Words exactly as briefed. Montserrat Bold headlines, Source Sans 3 small text.
  python3 social/regenerative-2026-10-15/build.py
    -> layers/, slides.html, slide-1..5.png, contact-sheet.png, regenerative-canva.pptx (import into Canva)
"""
import os, re, json, subprocess
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
NEW = os.path.join(ROOT, "img", "new-2026-10")
LAY = os.path.join(HERE, "layers"); os.makedirs(LAY, exist_ok=True)
W, H = 1080, 1350
DEEP, GREEN, CREAM = "#22371F", "#156826", "#F4F1EA"
def rgb(h): return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))


PHOTOS = os.environ.get("PHOTOS") == "1"  # design only for now: photo slides get an empty slot


def photo45(src, fx=0.5, fy=0.5):
    if not PHOTOS:
        return Image.new("RGB", (W, H), (196, 186, 168))
    """Crop to 4:5 and scale to 1080 x 1350. Refuses sources under 1080 px wide and any upscale."""
    im = Image.open(src).convert("RGB"); sw, sh = im.size
    assert sw >= 1080, f"{src} is under 1080 px wide"
    if sw / sh > W / H: ch, cw = sh, round(sh * W / H)
    else: cw, ch = sw, round(sw * H / W)
    assert cw >= W, f"{src}: 4:5 crop is {cw} px wide, would need upscaling"
    x, y = round((sw - cw) * fx), round((sh - ch) * fy)
    return im.crop((x, y, x + cw, y + ch)).resize((W, H), Image.LANCZOS)


SLIDES = []
class Slide:
    def __init__(self, n, bg=None, photo=None):
        self.n, self.items = n, []; SLIDES.append(self)
        im = photo if photo is not None else Image.new("RGB", (W, H), rgb(bg))
        self.image("background", im, 0, 0, jpg=True)

    def image(self, name, im, x, y, jpg=False):
        fn = f"s{self.n}-{len(self.items) + 1:02d}-{name}." + ("jpg" if jpg else "png")
        (im.convert("RGB").save(os.path.join(LAY, fn), quality=90) if jpg else im.save(os.path.join(LAY, fn)))
        self.items.append(dict(type="image", name=name, file="layers/" + fn, x=x, y=y, w=im.width, h=im.height))

    def band(self, name, y, h):
        self.image(name, Image.new("RGB", (W, h), rgb(DEEP)), 0, y, jpg=True)

    def text(self, name, runs, x, y, w, h, size, font, weight, lh):
        """runs: list of (text, colour, newline_after)."""
        self.items.append(dict(type="text", name=name, runs=runs, x=x, y=y, w=w, h=h, size=size, font=font, weight=weight, lh=lh))


logo = Image.open(os.path.join(ROOT, "img", "sfwlogo-240.png")).convert("RGBA")
logo = logo.resize((168, round(168 * logo.height / logo.width)), Image.LANCZOS)

SOILC = os.path.join(ROOT, "social", "cutouts-soil"); FLOW = os.path.join(ROOT, "social", "cutouts-flowers")
def cut(path, w=None, h=None, rot=0):
    im = Image.open(path).convert("RGBA")
    im = im.resize((w, round(im.height * w / im.width)) if w else (round(im.width * h / im.height), h), Image.LANCZOS)
    return im.rotate(rot, expand=True, resample=Image.BICUBIC) if rot else im
def put(s, name, im, cx, cy): s.image(name, im, round(cx - im.width / 2), round(cy - im.height / 2))

# 1. living soil, band at the top (design only: soil cutouts instead of a photo)
s = Slide(1, bg=CREAM)
put(s, "dark-soil", cut(os.path.join(SOILC, "dark-soil-clump.png"), w=720), 540, 860)
put(s, "earthworm", cut(os.path.join(SOILC, "earthworm.png"), h=360, rot=-30), 300, 1090)
put(s, "fern-fiddlehead", cut(os.path.join(SOILC, "fern-fiddlehead.png"), w=300), 780, 620)
put(s, "mushrooms", cut(os.path.join(SOILC, "mushrooms.png"), w=300), 860, 1160)
s.band("band", 0, 330)
s.text("headline", [("What does Regenerative actually mean?", CREAM, False)], 72, 64, 936, 160, 64, "Montserrat", 700, 1.15)
s.text("small-line", [("And who gets to decide?", CREAM, False)], 72, 236, 936, 56, 42, "Source Sans 3", 400, 1.3)

# 2 to 4. cream, the part before "or" in Deep Green, from "or" on in Food Web Green; same layout on all three
for n, before, after in [(2, "A destination", "or a direction?"),
                         (3, "A certified status", "or a way of thinking?"),
                         (4, "A premium product,", "or a standard everyone who works the land can reach?")]:
    s = Slide(n, bg=CREAM)
    s.image("logo", logo, W - 72 - logo.width, 64)
    s.text("question", [(before, DEEP, True), (after, GREEN, False)], 96, 420, 888, 760, 96, "Montserrat", 700, 1.1)

# 5. band at the bottom (design only: flowers from the four workshop regions instead of a photo)
s = Slide(5, bg=CREAM)
for nm, f, h_, cx, cy, r in [("poppy", "norfolk-common-poppy.png", 520, 160, 520, -6), ("marigold", "tamil-nadu-marigold.png", 420, 400, 560, 6),
                             ("paintbrush", "santa-fe-indian-paintbrush.png", 560, 640, 470, -4), ("heliconia", "costa-rica-heliconia.png", 600, 900, 460, 8)]:
    put(s, "flower-" + nm, cut(os.path.join(FLOW, f), h=h_, rot=r), cx, cy)
s.band("band", H - 440, 440)
s.text("headline", [("Living soil in the hands of every person who works the land.", CREAM, False)], 72, H - 440 + 64, 936, 230, 60, "Montserrat", 700, 1.15)
s.text("small-line", [("What does regenerative mean to you? Tell us in the comments.", CREAM, False)], 72, H - 440 + 300, 936, 110, 42, "Source Sans 3", 400, 1.3)

# ---------- outputs ----------
json.dump([dict(n=s.n, items=s.items) for s in SLIDES], open(os.path.join(HERE, "layout.json"), "w"), indent=1)
css = """@font-face{font-family:"Montserrat";src:url("../../fonts/montserrat-latin-variable.woff2");font-weight:100 900}
@font-face{font-family:"Source Sans 3";src:url("../../fonts/source-sans-3-latin-400-normal.woff2");font-weight:400}
*{margin:0;box-sizing:border-box}body{background:#555;display:grid;gap:40px;padding:40px;justify-content:center}
.slide{position:relative;width:1080px;height:1350px;overflow:hidden}.slide img{position:absolute;display:block}
.slide p{position:absolute}"""
out = [f"<!doctype html><html><head><meta charset='utf-8'><title>Regenerative carousel</title><style>{css}</style></head><body>"]
for s in SLIDES:
    out.append(f"<section class='slide' id='slide-{s.n}'>")
    for it in s.items:
        if it["type"] == "image":
            out.append(f"<img src='{it['file']}' style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;height:{it['h']}px' alt=''>")
        else:
            body = "".join(f"<span style='color:{c}'>{t}</span>" + ("<br>" if nl else "") for t, c, nl in it["runs"])
            out.append(f"<p style='left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;font-family:\"{it['font']}\";"
                       f"font-size:{it['size']}px;font-weight:{it['weight']};line-height:{it['lh']}'>{body}</p>")
    out.append("</section>")
out.append("</body></html>")
open(os.path.join(HERE, "slides.html"), "w").write("\n".join(out))
subprocess.run(["node", os.path.join(ROOT, "social", "turns-one-2026-10-16", "render.cjs"), HERE], check=True)

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
E = 914400 / 108
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(round(W * E)), Emu(round(H * E))
for s in SLIDES:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for it in s.items:
        if it["type"] == "image":
            sl.shapes.add_picture(os.path.join(HERE, it["file"]), Emu(round(it["x"] * E)), Emu(round(it["y"] * E)),
                                  Emu(round(it["w"] * E)), Emu(round(it["h"] * E))).name = it["name"]
    for it in s.items:
        if it["type"] != "text": continue
        tb = sl.shapes.add_textbox(Emu(round(it["x"] * E)), Emu(round(it["y"] * E)), Emu(round(it["w"] * E)), Emu(round(it["h"] * E)))
        tb.name = it["name"]; tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        for t, c, nl in it["runs"]:
            p.line_spacing = it["lh"]
            r = p.add_run(); r.text = t; f = r.font
            f.name = it["font"]; f.size = Pt(round(it["size"] * 72 / 108, 1)); f.bold = it["weight"] >= 600
            f.color.rgb = RGBColor(*rgb(c))
            if nl: p = tf.add_paragraph()
prs.save(os.path.join(HERE, "regenerative-canva.pptx"))
print("done:", len(SLIDES), "slides")
