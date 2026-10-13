"""'What does Regenerative actually mean?' single image, three visual concepts with real soil material,
gradients and grain. Headline and small line only.

  python3 social/regenerative-2026-10-15/concepts.py  -> concept-1/ .. concept-3/ (post.png, pptx) and concepts-compare.png
"""
import os
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import variants as v
from variants import Deck, W, H, DEEP, GREEN, LIGHT, CREAM, rgb, cut, logo_img, NEW, HEAD, SMALL1

HERE = os.path.dirname(os.path.abspath(__file__))
SOIL = "#4F3433"
MONT = os.path.join(HERE, "..", "open-house-2026-10-16", "fonts", "montserrat-latin-variable.ttf")


class Single(Deck):
    def __init__(self, key):
        super().__init__(key); self.dir = os.path.join(HERE, f"concept-{key}")
        self.lay = os.path.join(self.dir, "layers"); os.makedirs(self.lay, exist_ok=True)

    def write(self):
        super().write()
        os.replace(os.path.join(self.dir, "slide-1.png"), os.path.join(self.dir, "post.png"))
        if os.path.exists(os.path.join(self.dir, "contact-sheet.png")): os.remove(os.path.join(self.dir, "contact-sheet.png"))
        from pptx import Presentation
        src = os.path.join(self.dir, f"regenerative-option-{self.key}-canva.pptx")
        prs = Presentation(src)
        for sh in prs.slides[0].shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.name == "Source Sans 3": r.font.name = "Source Sans Pro"
        prs.save(os.path.join(self.dir, f"regenerative-concept-{self.key}-canva.pptx")); os.remove(src)

    def bg(self, s, im):
        s["items"].clear(); self.img(s, "background", im, 0, 0, jpg=True)


def grain(im, amount=9, seed=0):
    a = np.asarray(im.convert("RGB"), np.float32)
    a += np.random.default_rng(seed).normal(0, amount, a.shape[:2])[..., None]
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))


def cover(im, w, h, fx=.5, fy=.5):
    s = max(w / im.width, h / im.height); im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    x, y = round((im.width - w) * fx), round((im.height - h) * fy); return im.crop((x, y, x + w, y + h))


def radial(c_in, c_out, cx=.5, cy=.45, r=.9):
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    d = np.clip(np.hypot((xx / W - cx) * W / H, yy / H - cy) / r, 0, 1)[..., None]
    a, b = np.array(rgb(c_in), np.float32), np.array(rgb(c_out), np.float32)
    return Image.fromarray((a * (1 - d) + b * d).astype(np.uint8))


def build_concepts():
    # 1. Cracked to alive: bare, degraded earth at the top becomes living compost at the bottom
    d = Single("1"); s = d.slide(CREAM)
    dry = cover(Image.open(os.path.join(HERE, "textures", "dry-cracked-earth.png")).convert("RGB"), W, H)
    alive = cover(Image.open(os.path.join(NEW, "R5A_4007.jpg")).convert("RGB"), W, H, .5, .55)
    yy = np.mgrid[0:H, 0:W][0].astype(np.float32) / H
    noise = np.array(Image.fromarray((np.random.default_rng(3).random((H // 24, W // 24)) * 255).astype(np.uint8)).resize((W, H), Image.BICUBIC), np.float32) / 255
    m = np.clip((yy - 0.48 + (noise - .5) * .16) / 0.16, 0, 1)            # a ragged, soaking seam rather than a straight line
    mask = Image.fromarray((m * 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(6))
    base = Image.composite(alive, dry, mask)
    wash = np.clip(1 - yy / 0.42, 0, 1)[..., None] * 0.62
    base = Image.fromarray((np.asarray(base, np.float32) * (1 - wash) + np.array(rgb(CREAM), np.float32) * wash).astype(np.uint8))
    d.bg(s, grain(base, 6, 1))
    d.text(s, "headline", [(HEAD, SOIL, False)], 72, 110, 936, 340, 100, lh=1.04)
    d.text(s, "small-line", [(SMALL1, GREEN, False)], 76, 440, 900, 64, 52, font="Source Sans 3", weight=600)
    im = cut("fern-fiddlehead.png", w=260); d.img(s, "fern-fiddlehead", im, 800, 790, center=True)
    d.img(s, "logo", logo_img(110, disc=True), W - 72 - 138, H - 72 - 138)
    d.write()

    # 2. Under the microscope: a real student micrograph as a lens on a deep green radial gradient
    d = Single("2"); s = d.slide(DEEP)
    d.bg(s, grain(radial("#2E4A29", "#142112", .5, .62, 1.0), 8, 2))
    micro = Image.open(os.path.join(NEW, "Testate amoeba, fungal hypha, actinobacteria, Lactobacilli, 40x obj, Luke Hardy.jpg")).convert("RGB")
    D = 700; lens = cover(micro, D, D, .35, .45).convert("RGBA")
    mk = Image.new("L", (D, D), 0); ImageDraw.Draw(mk).ellipse((0, 0, D - 1, D - 1), fill=255); lens.putalpha(mk)
    ring = Image.new("RGBA", (D + 28, D + 28), (0, 0, 0, 0)); ImageDraw.Draw(ring).ellipse((0, 0, D + 27, D + 27), fill=rgb(CREAM) + (255,))
    ring.alpha_composite(lens, (14, 14)); glow = Image.new("RGBA", (D + 200, D + 200), (0, 0, 0, 0))
    ImageDraw.Draw(glow).ellipse((60, 60, D + 140, D + 140), fill=rgb(LIGHT) + (90,)); glow = glow.filter(ImageFilter.GaussianBlur(40))
    glow.alpha_composite(ring, (86, 86))
    d.img(s, "microscope-lens", glow, 540, 880, center=True)
    d.text(s, "headline", [(HEAD, CREAM, False)], 72, 90, 936, 300, 88, lh=1.04)
    d.text(s, "small-line", [(SMALL1, LIGHT, False)], 76, 390, 900, 64, 50, font="Source Sans 3", weight=600)
    d.text(s, "credit", [("Photo: Luke Hardy, 40x", CREAM, False)], 76, H - 64, 600, 30, 24, font="Source Sans 3", weight=400)
    d.img(s, "logo", logo_img(100, disc=True), W - 72 - 125, H - 60 - 125)
    d.write()

    # 3. Soil in the letters: "Regenerative" cut from a real compost photo, on a warm cream gradient with grain
    d = Single("3"); s = d.slide(CREAM)
    d.bg(s, grain(radial(CREAM, "#D8C9AC", .5, .42, 1.1), 7, 4))
    font = ImageFont.truetype(MONT, 196); font.set_variation_by_axes([900])
    word = "Regenerative"; bb = font.getbbox(word); tw = bb[2] - bb[0]
    font = ImageFont.truetype(MONT, int(196 * (W - 110) / tw)); font.set_variation_by_axes([900])
    bb = font.getbbox(word); tw, th = bb[2] - bb[0], bb[3] - bb[1]
    tm = Image.new("L", (tw + 20, th + 20), 0); ImageDraw.Draw(tm).text((10 - bb[0], 10 - bb[1]), word, font=font, fill=255)
    soil = cover(Image.open(os.path.join(NEW, "card-roots-in-soil-capri.jpg")).convert("RGB"), tm.width, tm.height, .45, .35)
    wimg = soil.convert("RGBA"); wimg.putalpha(tm)
    d.text(s, "what-does", [("What does", SOIL, False)], 58, 360, 936, 110, 96, lh=1.0)
    d.img(s, "word-regenerative-soil", wimg, 540, 545, center=True)
    d.text(s, "actually-mean", [("actually mean?", SOIL, False)], 58, 625, 936, 110, 96, lh=1.0)
    d.text(s, "small-line", [(SMALL1, GREEN, False)], 62, 760, 900, 64, 52, font="Source Sans 3", weight=600)
    d.img(s, "logo", logo_img(150), 72, H - 72 - 130)
    im = cut("fern-fiddlehead.png", w=240); d.img(s, "fern-fiddlehead", im, 760, 1170, center=True)
    im = cut("mushrooms.png", w=240); d.img(s, "mushrooms", im, 950, 1190, center=True)
    d.write()

    tw_, th_, g_ = 480, 600, 30
    from PIL import ImageFont as F
    cmp_ = Image.new("RGB", (g_ + 3 * (tw_ + g_), th_ + 100), (31, 31, 31)); dr = ImageDraw.Draw(cmp_)
    f = F.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 26)
    for i, (k, n) in enumerate([("1", "Cracked to alive"), ("2", "Under the microscope"), ("3", "Soil in the letters")]):
        x = g_ + i * (tw_ + g_); dr.text((x, 18), n, font=f, fill=(240, 240, 240))
        cmp_.paste(Image.open(os.path.join(HERE, f"concept-{k}", "post.png")).convert("RGB").resize((tw_, th_), Image.LANCZOS), (x, 60))
    cmp_.save(os.path.join(HERE, "concepts-compare.png")); print("done")


if __name__ == "__main__":
    build_concepts()
